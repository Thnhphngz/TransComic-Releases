"""Opt-in real local batch, pause/restart/resume on generated fixtures only."""
import argparse
import json
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--models',required=True);parser.add_argument('--output',required=True);args=parser.parse_args()
    from docx import Document
    from transcomic.tts.storage import TtsStore
    from transcomic.tts.documents import read_document
    from transcomic.tts.models import ModelManager
    from transcomic.tts.output import Outputs
    from transcomic.tts.scheduler import Scheduler
    from transcomic.tts.audio import validate
    root=Path(args.output).absolute();root.mkdir(exist_ok=False,parents=True)
    source=root/'fixtures';source.mkdir();audio=root/'Audio';audio.mkdir();manager=ModelManager(args.models)
    text=source/'Chapter 1.txt';text.write_text('Good morning.\nA quiet day begins.',encoding='utf8')
    word=source/'Chapter 2.docx';document=Document();p=document.add_paragraph('The sky is clear.');p.add_run().add_break();p.add_run('The birds are singing.');document.save(word)
    store=TtsStore(root/'tts.sqlite3');run=store.draft()
    for path in (text,word):store.add_document(run,read_document(path))
    store.add_document(run,dict(name='Pasted text',text='Thank you for listening.',line_count=1,path=None))
    config=dict(device='CPU',voice='af_heart',speed=1,language='en-us',variant='standard',format='wav')
    outputs=Outputs(store);plan=outputs.plan(run,'custom',str(audio),'wav');outputs.confirm(run,config,plan)
    def notify(kind,value):
        if kind=='segment' and store.segment(value)['state']=='running':scheduler.pause()
    scheduler=Scheduler(store,manager,notify);scheduler.start(run);scheduler.thread.join(120)
    assert not scheduler.busy and store.progress(run)=={'pending':4,'success':1}
    first=next(i for i in store.items(run) if i['state']=='success');saved=first['sha256'];store.close()
    store=TtsStore(root/'tts.sqlite3');assert store.run(run)['state']=='paused'
    scheduler=Scheduler(store,manager);scheduler.start(run);scheduler.thread.join(180)
    assert not scheduler.busy and store.run(run)['state']=='completed'
    assert store.segment(first['id'])['sha256']==saved
    docs=store.docs(run)
    for item in store.items(run):
        doc=store.doc(item['document_id']);assert validate(Path(doc['folder'])/item['filename'])['sha256']==item['sha256']
    report=dict(run=run,documents=[dict(name=d['name'],folder=d['folder']) for d in docs],counts=store.progress(run),pause_restart_resume=True,first_audio_preserved=True)
    (root/'report.json').write_text(json.dumps(report,indent=2),encoding='utf8');print(json.dumps(report,indent=2));store.close()


if __name__=='__main__':main()
