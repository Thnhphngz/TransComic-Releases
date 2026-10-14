"""Native run_transcomic.py Stage 2 smoke; isolated storage, mock HTTP only."""
import asyncio
import json
import os
import runpy
import time
from pathlib import Path
try:
    from scripts.test_runtime import TemporaryDirectory
except ModuleNotFoundError:
    from test_runtime import TemporaryDirectory
import httpx
from docx import Document
from PySide6.QtCore import QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox
from transcomic.bulk import plan_exports, write_exports
from transcomic.gemini import GeminiProvider
from transcomic.ui.main import MainWindow
import transcomic.ui.stage2 as controls
import transcomic.app as application

OUT = Path('artifacts/gui/stage2')


def wait_until(predicate, timeout=20):
    deadline = time.monotonic()+timeout
    while not predicate():
        QApplication.instance().processEvents(); QTest.qWait(10)
        if time.monotonic() > deadline: raise AssertionError('GUI smoke deadline')


def exercise(window):
    app, store = QApplication.instance(), window.store
    assert window.doc_id is None and not window.model.all_rows
    calls = []
    async def handler(request):
        calls.append(request)
        await asyncio.sleep(.08)
        payload = json.loads(json.loads(request.content)['contents'][0]['parts'][0]['text'])
        return httpx.Response(200, json={'candidates': [{'content': {'parts': [{'text': json.dumps({'lines': [{'id': r['id'], 'text': 'Mock '+r['text']} for r in payload['lines']]})}]}, 'finishReason': 'STOP'}], 'usageMetadata': {'promptTokenCount': 80, 'candidatesTokenCount': 30, 'totalTokenCount': 110}})
    original_provider = controls.GeminiProvider
    original_info = QMessageBox.information
    controls.GeminiProvider = lambda keys, model, **kwargs: GeminiProvider(keys, model, transport=httpx.MockTransport(handler), **kwargs)
    QMessageBox.information = lambda *args: QMessageBox.Ok
    try:
        source = store.path.parent/'input'; source.mkdir(); (source/'nested').mkdir()
        for folder in (source, source/'nested'):
            (folder/'Chapter 2.txt').write_text('\n'.join('Source line '+str(i) for i in range(24)), encoding='utf-8')
        word = Document(); word.add_paragraph('Word paragraph'); word.add_table(rows=1, cols=2).cell(0, 0).text = 'Word table'
        word.save(source/'Chapter 10.docx')
        originals = {str(p): p.read_bytes() for p in source.rglob('*') if p.is_file()}
        window.source_lang.setCurrentText('Tiếng Việt'); window.target_lang.setCurrentText('Tiếng Anh')
        window.import_paths([source, source/'Chapter 2.txt'])
        wait_until(lambda: not window.import_busy)
        assert len(window.workspace_ids) == 3
        store.add_gemini_keys('smoke-private-key', 'smoke-project')
        store.set_setting('gemini_models', json.dumps({'models': [{'name': 'models/mock-text', 'displayName': 'Mock text'}]}))
        store.set_setting('gemini_model', 'models/mock-text')
        window.provider_selector.setCurrentIndex(1)
        for name, value in {'rpm': 1000, 'tpm': 1000000, 'rpd': 1000, 'concurrency': 3}.items(): window.gemini_limit_controls[name].setValue(value)
        window.change_provider(); window.concurrency.setValue(3); window.gemini_verified.setChecked(True)
        window.extra_targets = ['Tiếng Indonesia']
        window.start_translation()
        wait_until(lambda: len(calls) >= 3)
        window.pause_or_resume()
        wait_until(lambda: window.scheduler.active == 0)
        assert window.scheduler._paused.is_set()
        doc_id = window.workspace_ids[0]
        window.open_document(doc_id)
        row = store.lines(doc_id)[0]
        window.model.setData(window.model.index(0, 1), 'Manual survives API')
        window.open_document(window.workspace_ids[1]); window.open_document(doc_id)
        assert store.line(row['id'])['translation'] == 'Manual survives API'
        window.grab().save(str(OUT/'light-paused-manuscript.png'))
        window.pause_or_resume()
        wait_until(lambda: not window.scheduler._thread.is_alive())
        app.processEvents(); window.refresh()
        assert len(window.workspace_ids) == 6
        assert store.line(row['id'])['translation'] == 'Manual survives API'
        assert all(r['translation'] for d in window.workspace_ids for r in store.lines(d) if r['source'].strip())
        window.pages.setCurrentIndex(1); window.refresh_stage2(); QTest.qWait(80)
        window.grab().save(str(OUT/'light-queue.png'))
        exports = write_exports(store, plan_exports(store.export_snapshot(window.workspace_ids), store.path.parent/'export'))
        assert len(exports) == 6 and all(e['state'] == 'exported' for e in exports)
        assert all(Path(p).read_bytes() == raw for p, raw in originals.items())
        window.toggle_theme(); window.resize(1000, 700); window.pages.setCurrentIndex(0); QTest.qWait(100)
        print('Window sizing:', window.width(), window.height(), window.minimumSizeHint().width(), [(i, window.pages.widget(i).minimumSizeHint().width()) for i in range(window.pages.count())])
        window.grab().save(str(OUT/'dark-manuscript-1000.png'))
        window.pages.setCurrentIndex(2); QTest.qWait(100)
        window.grab().save(str(OUT/'dark-api-1000.png'))
        window.pages.setCurrentIndex(4); window.console.reload(); QTest.qWait(80)
        window.grab().save(str(OUT/'dark-console-1000.png'))
        report = {'platform': app.platformName(), 'scale': os.environ.get('QT_SCALE_FACTOR'), 'startup_empty': True, 'imports': 3, 'workspaces': 6, 'targets': ['EN', 'ID'], 'concurrency': 3, 'mock_translation_requests': len(calls), 'real_api_requests': 0, 'exports': len(exports), 'manual_preserved': True, 'pause_resume': True, 'source_unchanged': True}
        (OUT/'report.json').write_text(json.dumps(report, indent=2), encoding='utf-8'); print(json.dumps(report))
    finally:
        controls.GeminiProvider = original_provider; QMessageBox.information = original_info
        if window.scheduler: window.scheduler.stop(); window.scheduler.join(5)
        window._exit_authorized = True; window.close(); app.quit()


class SmokeWindow(MainWindow):
    def __init__(self, store):
        super().__init__(store)
        def run():
            try: exercise(self)
            except Exception:
                import traceback
                traceback.print_exc(); self._exit_authorized = True; self.close(); QApplication.instance().exit(1)
        QTimer.singleShot(150, run)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory() as temporary:
        previous = os.environ.get('LOCALAPPDATA'); os.environ['LOCALAPPDATA'] = temporary
        application.MainWindow = SmokeWindow
        try: runpy.run_path('run_transcomic.py', run_name='__main__')
        finally:
            if previous is None: os.environ.pop('LOCALAPPDATA', None)
            else: os.environ['LOCALAPPDATA'] = previous


if __name__ == '__main__': main()
