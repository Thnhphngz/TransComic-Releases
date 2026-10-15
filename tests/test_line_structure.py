import asyncio
import copy
import io
import json
from pathlib import Path
import pytest
from docx import Document
from transcomic.documents import decode_txt, render_export, split_lines, ImportedText, TextLine
from transcomic.docx_document import inspect_docx
from transcomic.storage import Store
from transcomic.providers import KiraProvider, parse_line_response
from transcomic.gemini import GeminiProvider
from transcomic.line_structure import validate_snapshot
from transcomic.bulk import plan_exports, write_exports


@pytest.fixture
def store(tmp_path):
    s=Store(tmp_path/'state.db');yield s;s.close()


def imported(store,path,raw=b'Alpha\nBeta\nGamma'):
    path.write_bytes(raw)
    return store.import_document(path,decode_txt(raw),'vi','en')[0]


def snapshot(store,doc): return store.export_snapshot([doc])[0]


def complete(store,doc):
    for row in store.lines(doc):
        if row['state']!='skip':store.edit_line(row['id'],'translation','Hello '+str(row['ordinal']))


@pytest.mark.parametrize('case',['valid','missing','extra','duplicate','unknown','reverse','newline','bad_type'])
def test_ten_ids_response_structure_and_mapping(case):
    entries=[{'id':str(i),'text':f'Result {i}'} for i in range(10)]
    if case=='missing':entries.pop()
    if case=='extra':entries.append({'id':'extra','text':'Extra'})
    if case=='duplicate':entries[-1]['id']='0'
    if case=='unknown':entries[-1]['id']='foreign'
    if case=='reverse':entries.reverse()
    if case=='newline':entries[-1]['text']='Merged\nSplit'
    if case=='bad_type':entries[-1]['id']=[]
    parsed=parse_line_response(json.dumps({'lines':entries}),set(map(str,range(10))))
    assert bool(parsed.error)==(case not in ('valid','reverse'))
    assert all(value==f'Result {key}' for key,value in parsed.values.items())
    if case in ('valid','reverse'):assert len(parsed.values)==10
    elif case=='extra':assert len(parsed.values)==10 and parsed.diagnostics['structure_errors']
    elif case=='duplicate':assert '0' not in parsed.values and '9' not in parsed.values


@pytest.mark.parametrize('raw',[b'',b'\nAlpha\n\nBeta\n\n',b'\r\nAlpha\r\n\r\n',b'\rAlpha\r\r',b'\nA\r\nB\r \nC',b'Alpha',b'Alpha\n'])
def test_txt_physical_units_and_empty_tail(store,tmp_path,raw):
    doc=imported(store,tmp_path/'a.txt',raw);complete(store,doc)
    d,rows=snapshot(store,doc);report=validate_snapshot(d,rows)
    assert not report.structural and report.total==len(decode_txt(raw).lines)
    result=write_exports(store,plan_exports([(d,rows)],tmp_path/'out'))[0]
    assert result['state']=='exported'
    assert len(split_lines(Path(result['path']).read_bytes().decode('utf8')))==report.total
    assert [r['ending'] for r in rows]==[u.ending for u in decode_txt(raw).lines]
    if not report.required:assert report.state=='incomplete'


@pytest.mark.parametrize('case',['missing','extra','duplicate_id','owner','ordinal','order','ending','source','newline','cr','number'])
def test_snapshot_corruption_cannot_pass_equal_counts(store,tmp_path,case):
    doc=imported(store,tmp_path/'a.txt');complete(store,doc)
    d,rows=snapshot(store,doc)
    if case=='missing':rows.pop()
    if case=='extra':rows.append({**rows[-1],'id':'extra','ordinal':3})
    if case=='duplicate_id':rows[-1]['id']=rows[0]['id']
    if case=='owner':rows[-1]['document_id']='foreign'
    if case=='ordinal':rows[-1]['ordinal']=0
    if case=='order':rows.reverse()
    if case=='ending':rows[0]['ending']=''
    if case=='source':rows[0]['source']='Other source'
    if case=='newline':rows[0]['translation']='One\nTwo'
    if case=='cr':rows[0]['translation']='One\rTwo'
    if case=='number':rows[0]['translation']=123
    report=validate_snapshot(d,rows)
    assert report.state=='structure_error' and report.structural
    target=tmp_path/'out.txt';target.write_bytes(b'Keep existing')
    before=set(tmp_path.iterdir())
    result=write_exports(store,[{'document':d,'rows':rows,'target':target}],existing='overwrite',allow_incomplete=True)[0]
    assert result['state']=='blocked_structure' and target.read_bytes()==b'Keep existing'
    assert set(tmp_path.iterdir())==before


def test_unfinished_states_are_neutral_and_source_edits_keep_mapping(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt',b'Alpha\nBeta\nGamma\nDelta\n\n123')
    rows=store.lines(doc);store.edit_line(rows[0]['id'],'translation','Hello.')
    with store.transaction():
        store.db.execute("UPDATE lines SET state='processing' WHERE id=?",(rows[1]['id'],))
        store.db.execute("UPDATE lines SET state='error',error='HTTP 504' WHERE id=?",(rows[2]['id'],))
    r=validate_snapshot(*snapshot(store,doc))
    assert r.state=='incomplete' and (r.total,r.required,r.translated,r.processing,r.pending,r.errors,r.skipped)==(6,4,1,1,1,1,2)
    store.edit_line(rows[3]['id'],'source','Changed source')
    assert not validate_snapshot(*snapshot(store,doc)).structural


def test_old_edited_source_clone_is_readable_without_migration(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');row=store.lines(doc)[0]
    store.edit_line(row['id'],'source','Edited source')
    variant=store.workspace_for_target(doc,'id');complete(store,variant)
    store.edit_line(row['id'],'source','Parent changed again after cloning')
    schema=store.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]
    before=store.db.total_changes;r=validate_snapshot(*snapshot(store,variant))
    assert r.state=='complete' and not r.structural and store.db.total_changes==before
    assert store.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]==schema


def test_docx_blank_paragraphs_inline_breaks_and_soft_wrap(store,tmp_path):
    obj=Document();obj.add_paragraph('Alpha '*300);obj.add_paragraph('')
    p=obj.add_paragraph('Beta');p.add_run().add_break();p.add_run('Gamma')
    data=io.BytesIO();obj.save(data);raw=data.getvalue();parsed=inspect_docx(raw)
    doc=store.import_document(tmp_path/'a.docx',ImportedText(tuple(TextLine(u.text,'') for u in parsed.units),'docx',raw),'vi','en','docx',[u.locator for u in parsed.units])[0]
    complete(store,doc);d,rows=snapshot(store,doc)
    assert len(rows)==4 and validate_snapshot(d,rows).state=='complete'
    plan=plan_exports([(d,rows)],tmp_path/'out');result=write_exports(store,plan)[0]
    assert result['state']=='exported' and len(inspect_docx(Path(result['path']).read_bytes()).units)==4
    rows[1]['locator']=rows[0]['locator'];assert validate_snapshot(d,rows).structural
    with store.transaction():store.db.execute('DELETE FROM document_units WHERE line_id=?',(rows[1]['id'],))
    assert len(store.units(doc))==4 and validate_snapshot(*snapshot(store,doc)).state=='structure_error'


def test_batch_evidence_and_manual_correction(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');job=store.create_job(doc,'mock','vi','en',1)
    batch=store.batches(job)[0];payload=store.claim_batch(batch['id'])
    parsed=parse_line_response(json.dumps({'lines':[{'id':payload[0]['id'],'text':'Hello'}]}),{r['id'] for r in payload})
    store.finish_batch(batch['id'],parsed.values,parsed.error,diagnostics=parsed.diagnostics)
    r=validate_snapshot(*snapshot(store,doc))
    assert r.state=='structure_error' and len(r.structural)==2 and r.translated==1
    for row in store.lines(doc):
        if not row['translation']:store.edit_line(row['id'],'translation','Corrected')
    assert validate_snapshot(*snapshot(store,doc)).state=='complete'


def test_extra_api_id_retains_valid_results_but_blocks_structure(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');job=store.create_job(doc,'mock','vi','en',1)
    batch=store.batches(job)[0];payload=store.claim_batch(batch['id'])
    entries=[{'id':p['id'],'text':'Hello'} for p in payload]+[{'id':'foreign','text':'Ignored'}]
    parsed=parse_line_response(json.dumps({'lines':entries}),{p['id'] for p in payload})
    store.finish_batch(batch['id'],parsed.values,parsed.error,diagnostics=parsed.diagnostics)
    assert all(r['translation']=='Hello' for r in store.lines(doc))
    assert validate_snapshot(*snapshot(store,doc)).state=='structure_error'


def test_storage_rejects_multiline_before_normalization_and_foreign_payload(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');foreign=imported(store,tmp_path/'other.txt',b'Other')
    job=store.create_job(doc,'mock','vi','en',1);batch=store.batches(job)[0];payload=store.claim_batch(batch['id'])
    store.finish_batch(batch['id'],{payload[0]['id']:'One\nTwo',payload[1]['id']:'Hello'})
    assert not store.line(payload[0]['id'])['translation'] and store.line(payload[1]['id'])['translation']=='Hello'
    other=store.lines(foreign)[0]
    with store.transaction():store.db.execute("UPDATE batches SET state='inflight',payload=? WHERE id=?",(json.dumps([other]),batch['id']))
    store.finish_batch(batch['id'],{other['id']:'Must not cross'})
    assert not store.line(other['id'])['translation']


def test_partial_export_default_blocks_draft_explicit_and_bulk_isolation(store,tmp_path):
    good=imported(store,tmp_path/'good.txt');bad=imported(store,tmp_path/'bad.txt');complete(store,good);complete(store,bad)
    docs=store.export_snapshot([good,bad]);docs[1][1][0]['translation']='Bad\nSplit'
    results=write_exports(store,plan_exports(docs,tmp_path/'out'))
    assert [r['state'] for r in results].count('exported')==1
    assert next(r for r in results if r['document_id']==bad)['state']=='blocked_structure'
    incomplete=imported(store,tmp_path/'incomplete.txt')
    plan=plan_exports(store.export_snapshot([incomplete]),tmp_path/'draft')
    assert write_exports(store,plan)[0]['state']=='error' and not plan[0]['target'].exists()
    # Existing language guard still applies; use matching source/target to test draft mechanics.
    with store.transaction():store.db.execute("UPDATE documents SET source_lang='en' WHERE id=?",(incomplete,))
    plan=plan_exports(store.export_snapshot([incomplete]),tmp_path/'draft')
    assert write_exports(store,plan,allow_incomplete=True)[0]['state']=='exported'
    assert 'INCOMPLETE' in plan[0]['target'].name


@pytest.mark.parametrize('provider',['kira','gemini'])
def test_duplicate_request_ids_rejected_before_http(provider):
    adapter=KiraProvider('fake') if provider=='kira' else GeminiProvider([{'id':'fake','secret':'fake','project':''}],'models/fake')
    with pytest.raises(ValueError):asyncio.run(adapter.translate([{'id':'same','source':'A'},{'id':'same','source':'B'}],'vi','en'))


@pytest.mark.parametrize('entries',[
    [{'id':'a','text':'Good'},{'id':'a','text':None}],
    [{'id':'a','text':None},{'id':'a','text':'Good'}],
])
def test_invalid_duplicate_cannot_leave_an_ambiguous_value(entries):
    parsed=parse_line_response(json.dumps({'lines':entries}),{'a'})
    assert parsed.error and not parsed.values


def test_export_keeps_read_snapshot_and_never_blocks_for_period_alone(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');complete(store,doc)
    first=store.lines(doc)[0];store.edit_line(first['id'],'translation','Kept...')
    captured=store.export_snapshot([doc]);plan=plan_exports(captured,tmp_path/'out')
    store.edit_line(first['id'],'translation','Live edit')
    before=store.line(first['id'])
    result=write_exports(store,plan)[0]
    assert result['state']=='exported' and Path(result['path']).read_text(encoding='utf8').startswith('Kept...')
    assert store.line(first['id'])==before


def test_retry_success_supersedes_prior_missing_evidence(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');job=store.create_job(doc,'mock','vi','en',1)
    batch=store.batches(job)[0];payload=store.claim_batch(batch['id'])
    parsed=parse_line_response(json.dumps({'lines':[{'id':payload[0]['id'],'text':'Hello'}]}),{r['id'] for r in payload})
    store.finish_batch(batch['id'],parsed.values,parsed.error,diagnostics=parsed.diagnostics)
    assert validate_snapshot(*snapshot(store,doc)).structural
    # Offline provider completion of a new explicitly created job, no network/retry.
    ids=[r['id'] for r in store.lines(doc) if not r['translation']]
    newer=store.create_job(doc,'mock','vi','en',1,ids,allow_shared_document=True)
    batch=store.batches(newer)[0];payload=store.claim_batch(batch['id'])
    store.finish_batch(batch['id'],{r['id']:'Correct' for r in payload})
    assert validate_snapshot(*snapshot(store,doc)).state=='complete'


def test_legacy_database_reopen_validation_does_not_mutate_data(tmp_path):
    path=tmp_path/'old.db';s=Store(path);doc=imported(s,tmp_path/'a.txt');complete(s,doc)
    ids=[r['id'] for r in s.lines(doc)];s.close();s=Store(path)
    try:
        before=s.lines(doc);count=s.db.total_changes;schema=s.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]
        assert validate_snapshot(*snapshot(s,doc)).state=='complete'
        assert s.lines(doc)==before and [r['id'] for r in before]==ids and s.db.total_changes==count
        assert s.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]==schema
    finally:s.close()


def test_empty_success_response_is_missing_structure_not_done(store,tmp_path):
    doc=imported(store,tmp_path/'a.txt');job=store.create_job(doc,'mock','vi','en',1)
    batch=store.batches(job)[0];store.claim_batch(batch['id']);store.finish_batch(batch['id'],{})
    report=validate_snapshot(*snapshot(store,doc))
    assert report.state=='structure_error' and report.translated==0 and len(report.structural)==3
