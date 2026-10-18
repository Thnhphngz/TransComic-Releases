import asyncio
import json
import sqlite3
import threading
import time
import tracemalloc
from pathlib import Path

import httpx
import pytest
from transcomic.documents import decode_txt, ImportedText, TextLine
from transcomic.storage import Store
from transcomic.security import protect_key, unprotect_key
from transcomic.kira_pool import KiraPoolProvider
from transcomic.translation import Scheduler
from transcomic.line_structure import validate_snapshot


def wait_for(fn, timeout=8):
    end=time.monotonic()+timeout
    while not fn() and time.monotonic()<end:time.sleep(.01)
    assert fn()


def lines(req):return json.loads(json.loads(req.content)['messages'][1]['content'])['lines']


def success(req):
    return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({'lines':[
        {'id':r['id'],'text':'The translated story continues'} for r in reversed(lines(req))]})},'finish_reason':'stop'}],
        'usage':{'prompt_tokens':10,'completion_tokens':12,'total_tokens':22}})


def setup(tmp_path, keys=2, per_key=2, files=12, batches=3, **policy):
    store=Store(tmp_path/'state.db')
    store.add_kira_keys(' '.join(f'sk-synthetic-{i:04d}-not-real' for i in range(keys)))
    ids=[k['id'] for k in store.kira_keys()]
    for key in ids:store.set_kira_account_scope(key,ids[0])
    store.set_setting('concurrency',str(per_key))
    docs=[]
    for i in range(files):
        raw='\r\n'.join(f'Nguồn tổng hợp {j}' for j in range(3*(batches if i%2 else 1))).encode()
        doc=store.import_document(tmp_path/f'{i}.txt',decode_txt(raw),'vi','en')[0];docs.append(doc)
    run,jobs=store.create_run(docs,['en'],provider='kira',model='mock',source_lang='vi',concurrency=per_key,config={'kira_multi_key':True,'kira_recovery_version':1,**policy})
    return store,docs,run,jobs


def start(store,run,jobs,handler,**options):
    provider=KiraPoolProvider(store,'mock',transport=httpx.MockTransport(handler),
        account_limit=options.pop('account_limit',150),rpm=options.pop('rpm',100000),tpm=options.pop('tpm',1000000000),**options)
    scheduler=Scheduler(store,provider);scheduler.start_many(jobs,run)
    return scheduler,provider


@pytest.mark.parametrize('count',[1,2,3,5])
def test_key_management_dpapi_duplicate_toggle_delete_restart(tmp_path,count):
    s=Store(tmp_path/'state.db');keys=[f'sk-only-fixture-{i:03d}' for i in range(count)]
    assert s.add_kira_keys(' '.join(keys))==(count,0)
    assert s.add_kira_keys(keys[0])==(0,1)
    metadata=s.kira_keys();assert all('secret' not in k and 'encrypted' not in k for k in metadata)
    first=metadata[0]['id'];s.change_kira_key(first,name='Production example',enabled=False)
    assert s.kira_keys()[0]['status']=='DISABLED'
    s.change_kira_key(first,enabled=True)
    assert s.kira_keys(decrypt=True)[0]['secret']==keys[0]
    s.close();s=Store(tmp_path/'state.db')
    assert s.kira_keys()[0]['id']==first and s.kira_keys()[0]['name']=='Production example'
    s.change_kira_key(first,delete=True);s.close();s=Store(tmp_path/'state.db')
    assert len(s.kira_keys())==count-1
    raw=Path(s.path).read_bytes()
    assert all(k.encode() not in raw for k in keys)
    s.close()


def test_legacy_ciphertext_atomic_migration_and_no_resurrection(tmp_path):
    path=tmp_path/'old.db';s=Store(path);blob=protect_key('sk-legacy-synthetic')
    with s.transaction():
        s.db.execute("DELETE FROM settings WHERE key='kira_keys_migrated'")
        s.db.execute("INSERT INTO secrets VALUES ('kira',?)",(blob,))
    original=s.db.total_changes
    with pytest.raises(RuntimeError):
        with s.transaction():s.migrate_kira_key();raise RuntimeError('simulate power failure')
    assert not s.kira_keys()
    s.close();s=Store(path)
    assert unprotect_key(s.load_key())=='sk-legacy-synthetic'
    assert s.kira_keys()[0]['id']=='legacy-kira'
    s.migrate_kira_key();assert len(s.kira_keys())==1
    assert s.add_kira_keys('sk-legacy-synthetic')==(0,1)
    s.change_kira_key('legacy-kira',delete=True);s.close();s=Store(path)
    assert s.load_key() is None and not s.kira_keys()
    s.close()


def test_key_add_rollback_on_dpapi_failure_and_no_plaintext_fallback(tmp_path,monkeypatch):
    s=Store(tmp_path/'s.db');s.add_kira_keys('sk-initial-example')
    import transcomic.security as security
    def fail(_):raise OSError('synthetic DPAPI failure')
    monkeypatch.setattr(security,'protect_key',fail)
    with pytest.raises(OSError):s.add_kira_keys('sk-next-example')
    assert len(s.kira_keys())==1;s.close()


@pytest.mark.parametrize('keys,cap',[(1,30),(2,30),(3,30),(5,30),(3,10)])
def test_mock_load_fair_slots_mapping_and_single_os_worker(tmp_path,keys,cap):
    s,docs,run,jobs=setup(tmp_path,keys,cap,files=25,batches=32)
    release=threading.Event()
    seen=[];busy={};peaks={};active_files=set();peak_files=0;os_threads=set()
    job_of={r['id']:j for j in jobs for b in s.batches(j) for r in b['payload']}
    source_of={j:s.document(s.job(j)['document_id'])['source_instance'] for j in jobs}
    async def handler(req):
        nonlocal peak_files
        key=req.headers['authorization'];batch=tuple(r['id'] for r in lines(req))
        assert batch not in seen;seen.append(batch)
        busy[key]=busy.get(key,0)+1;peaks[key]=max(peaks.get(key,0),busy[key])
        assert busy[key]<=cap
        os_threads.add(threading.get_ident())
        while not release.is_set():await asyncio.sleep(.005)
        await asyncio.sleep(.02 if key.endswith('0000-not-real') else .005)
        busy[key]-=1
        return success(req)
    tracemalloc.start()
    began=time.monotonic();scheduler,pool=start(s,run,jobs,handler)
    wait_for(lambda:scheduler.active==keys*cap,timeout=15)
    release.set()
    observations=[]
    while scheduler._thread.is_alive() and time.monotonic()-began<35:
        observations.append((scheduler.active,getattr(scheduler,'active_files',0)))
        time.sleep(.002)
    scheduler.join(1)
    assert not scheduler._thread.is_alive()
    assert s.run(run)['state']=='finished'
    assert max(a for a,f in observations)<=keys*cap and max(f for a,f in observations)<=10
    assert sum(s['peak'] for s in pool.stats.values())==keys*cap
    assert len(os_threads)==1 and all(p<=cap for p in peaks.values())
    assert len(seen)==sum(len(s.batches(j)) for j in jobs)
    assert len(peaks)==keys
    for doc in docs:
        snapshot,rows=s.export_snapshot([doc])[0];assert validate_snapshot(snapshot,rows).state=='complete'
    assert all(len(s.attempts(b['id']))==1 for j in jobs for b in s.batches(j))
    assert all(len(s.batches(j)[0]['payload'][0].get('context',[]))<=8 for j in jobs)
    assert all(not s.job(j)['retry_504'] and s.job(j)['auto_retries']==2 for j in jobs)
    evidence=Path('artifacts/kira-multikey');evidence.mkdir(parents=True,exist_ok=True)
    elapsed=time.monotonic()-began
    current_bytes,peak_bytes=tracemalloc.get_traced_memory();tracemalloc.stop()
    (evidence/f'load-{keys}x{cap}.json').write_text(json.dumps({'keys':keys,'per_key':cap,
        'theoretical':keys*cap,'peak_busy':max(a for a,f in observations),'peak_active_files':max(f for a,f in observations),
        'requests':len(seen),'seconds':elapsed,'lines_per_minute':sum(len(b) for b in seen)*60/elapsed,
        'per_key_stats':pool.snapshot(cap)['per_key'],'http_os_threads':len(os_threads),
        'mock_only':True,'peak_traced_bytes':peak_bytes,
        'mean_dispatch_seconds':sum(scheduler.dispatch_latencies)/len(scheduler.dispatch_latencies)}),encoding='utf-8')
    s.close()


@pytest.mark.parametrize('failure',['401','403','quota','rate','504','timeout','connect','500'])
def test_error_isolation_no_retry_no_key_failover_same_batch(tmp_path,failure):
    s,docs,run,jobs=setup(tmp_path,2,1,files=4,batches=2,kira_recovery_version=0)
    key_ids=[k['id'] for k in s.kira_keys()];seen=[];failed=[]
    for key in key_ids:s.set_kira_account_scope(key,key)
    async def handler(req):
        batch=tuple(r['id'] for r in lines(req));seen.append(batch)
        if req.headers['authorization'].endswith('0000-not-real') and not failed:
            failed.append(batch)
            if failure=='timeout':raise httpx.ReadTimeout('fake',request=req)
            if failure=='connect':raise httpx.ConnectError('fake',request=req)
            if failure=='quota':return httpx.Response(429,json={'error':{'code':'insufficient_quota'}})
            if failure=='rate':return httpx.Response(429,headers={'retry-after':'1'},json={'error':{'code':'rate_limit'}})
            return httpx.Response(int(failure))
        await asyncio.sleep(.01);return success(req)
    scheduler,pool=start(s,run,jobs,handler)
    wait_for(lambda:failed)
    scheduler.join(10)
    assert not scheduler._thread.is_alive()
    assert seen.count(failed[0])==1
    assert scheduler.retry_count==0
    assert s.run(run)['state']!='finished'
    first=s.kira_keys()[0]['status']
    assert first in ({'401':{'INVALID'},'403':{'PERMISSION_DENIED'},'quota':{'QUOTA_EXHAUSTED'},'rate':{'READY','COOLDOWN'}}.get(failure,{'TEMPORARY_ERROR','READY'}))
    assert not scheduler._paused.is_set()
    if failure not in ('timeout','connect'):assert pool.stats[key_ids[1]]['requests']>=1
    assert all(len(s.attempts(b['id']))<=1 for j in jobs for b in s.batches(j))
    s.close()


def test_all_keys_unavailable_requires_explicit_start_without_false_pause(tmp_path):
    s,docs,run,jobs=setup(tmp_path,1,1,files=3,kira_recovery_version=0)
    calls=[]
    def handler(req):
        calls.append(tuple(r['id'] for r in lines(req)))
        return httpx.Response(401) if len(calls)==1 else success(req)
    scheduler,pool=start(s,run,jobs,handler)
    scheduler.join(3)
    assert not scheduler._thread.is_alive() and not scheduler._paused.is_set()
    assert len(calls)==1 and s.run(run)['state']=='needs_intervention'
    s.add_kira_keys('sk-added-synthetic');time.sleep(.1);assert len(calls)==1
    scheduler.start_many(jobs,run);scheduler.join(8)
    assert not scheduler._thread.is_alive() and len(calls)>1 and len(calls)==len(set(calls))
    s.close()


def test_pause_dynamic_disable_concurrency_stop_and_manual_edit_wins(tmp_path):
    s,docs,run,jobs=setup(tmp_path,2,2,files=12,batches=5)
    calls=[];entered=threading.Event();release=threading.Event()
    async def handler(req):
        calls.append((req.headers['authorization'],lines(req)));entered.set()
        while not release.is_set():await asyncio.sleep(.01)
        return success(req)
    scheduler,pool=start(s,run,jobs,handler)
    wait_for(lambda:len(calls)==4);scheduler.pause();count=len(calls)
    line_id=calls[0][1][0]['id'];s.edit_line(line_id,'translation','Manual edit wins')
    s.change_kira_key(s.kira_keys()[0]['id'],enabled=False)
    s.set_setting('concurrency','1');release.set();time.sleep(.1)
    assert len(calls)==count and s.line(line_id)['translation']=='Manual edit wins'
    scheduler.resume();wait_for(lambda:len(calls)>count);scheduler.stop();scheduler.join(5)
    assert not scheduler._thread.is_alive()
    assert s.line(line_id)['translation']=='Manual edit wins'
    assert all(not key.endswith('0000-not-real') for key,rows in calls[count:])
    s.close()


def test_claim_race_and_attempt_fence_recovery_idempotent(tmp_path):
    s,docs,run,jobs=setup(tmp_path,2,1,files=1)
    batch=s.batches(jobs[0])[0];claims=[]
    threads=[threading.Thread(target=lambda:claims.append(s.claim_batch(batch['id']))) for _ in range(2)]
    for t in threads:t.start()
    for t in threads:t.join()
    assert sum(bool(c) for c in claims)==1
    attempt=s.prepare_attempt(batch['id'],1)
    assert s.finish_batch(batch['id'],{r['id']:'Wrong' for r in batch['payload']},expected_attempt_id='stale')==[]
    s.close();s=Store(tmp_path/'state.db')
    assert s.batch(batch['id'])['state']=='unknown' and s.attempts(batch['id'])[0]['state']=='unknown'
    s.close();s=Store(tmp_path/'state.db')
    assert s.batch(batch['id'])['state']=='unknown'
    assert s.finish_batch(batch['id'],{r['id']:'Late' for r in batch['payload']},expected_attempt_id=attempt)==[]
    s.close()


def test_durable_account_budget_and_shared_cooldown(tmp_path):
    s,_,_,_=setup(tmp_path,3,2,files=1)
    pool=KiraPoolProvider(s,'mock',account_limit=2,rpm=2,tpm=100)
    a=pool.reserve(2,40);b=pool.reserve(2,40)
    assert a and b and pool.reserve(2,1) is None
    pool.release(a[0]);pool.release(b[0]);s.close()
    s=Store(tmp_path/'state.db');pool=KiraPoolProvider(s,'mock',rpm=2,tpm=100)
    assert pool.reserve(2,1) is None
    s.close()


@pytest.mark.parametrize('problem',['missing','extra','duplicate','newline'])
def test_partial_mapping_error_keeps_verified_data_and_blocks_export(tmp_path,problem):
    from transcomic.bulk import plan_exports,write_exports
    s,docs,run,jobs=setup(tmp_path,2,2,files=1,kira_recovery_version=0)
    def handler(req):
        rows=[{'id':r['id'],'text':'The story continues'} for r in lines(req)]
        if problem=='missing':rows.pop()
        if problem=='extra':rows.append({'id':'foreign-id','text':'Unmapped'})
        if problem=='duplicate':rows.append(dict(rows[0]))
        if problem=='newline':rows[0]['text']='Invalid\nSplit'
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({'lines':rows})}}]})
    scheduler,pool=start(s,run,jobs,handler);scheduler.join(5)
    assert scheduler.request_count==1 and scheduler.retry_count==0
    assert sum(v['errors'] for v in pool.stats.values())==1
    assert sum(v['successful_requests'] for v in pool.stats.values())==0
    assert sum(v['batches'] for v in pool.stats.values())==0
    assert sum(v['translated_lines'] for v in pool.stats.values())==(3 if problem=='extra' else 2)
    doc,rows=s.export_snapshot(docs)[0]
    report=validate_snapshot(doc,rows)
    if problem=='extra':
        # An out-of-batch ID was discarded. All three intended IDs were uniquely
        # verified and committed; retain the failed request in monitor history.
        assert not report.structural and report.state=='complete'
    else:assert report.structural and report.state=='structure_error'
    assert any(r['translation'] for r in rows)
    result=write_exports(s,plan_exports([(doc,rows)],tmp_path/'out'),existing='overwrite',allow_incomplete=True)
    if problem=='extra':assert result[0]['state']=='exported'
    else:assert result[0]['state']=='blocked_structure' and not (tmp_path/'out').exists()
    s.close()


def test_multiple_languages_order_and_docx_blank_units(tmp_path):
    from docx import Document
    from transcomic.docx_document import import_docx
    s=Store(tmp_path/'s.db');s.add_kira_keys('sk-fixture-one sk-fixture-two')
    path=tmp_path/'sample.docx';word=Document();word.add_paragraph('');word.add_paragraph('Nguồn một');word.add_paragraph('Nguồn hai');word.save(path)
    imported=import_docx(path)
    # Match the current DOCX importer API rather than visual Word lines.
    data=ImportedText(tuple(TextLine(u.text,'') for u in imported.units),'docx',imported.raw)
    doc=s.import_document(path,data,'vi','en',document_format='docx',locators=[u.locator for u in imported.units])[0]
    run,jobs=s.create_run([doc],['en','de'],provider='kira',model='mock',source_lang='vi',concurrency=2,config={'kira_multi_key':True,'kira_recovery_version':1})
    order=[]
    def handler(req):
        content=json.loads(req.content)['messages'][0]['content'];order.append(content)
        return success(req)
    scheduler,pool=start(s,run,jobs,handler);scheduler.join(6)
    assert not scheduler._thread.is_alive() and len(order)==2
    assert 'English' in order[0] and 'German' in order[1]
    assert all(len(s.lines(s.job(j)['document_id']))==3 for j in jobs)
    assert all(validate_snapshot(*s.export_snapshot([s.job(j)['document_id']])[0]).state=='complete' for j in jobs)
    s.close()


def test_safe_exit_cancellation_fences_late_results(tmp_path):
    s,docs,run,jobs=setup(tmp_path,2,2,files=3)
    async def handler(req):await asyncio.sleep(30);return success(req)
    scheduler,pool=start(s,run,jobs,handler)
    wait_for(lambda:scheduler.active_http==4)
    scheduler.cancel_inflight();scheduler.join(5)
    assert not scheduler._thread.is_alive() and not scheduler.active
    unknown=[b for j in jobs for b in s.batches(j) if b['state']=='unknown']
    assert len(unknown)==4
    assert all(not r['translation'] for d in docs for r in s.lines(d))
    s.close()


def test_no_keys_and_secret_redaction(tmp_path):
    s,docs,run,jobs=setup(tmp_path,2,1,files=1,kira_recovery_version=0)
    secrets=[k['secret'] for k in s.kira_keys(decrypt=True)]
    def handler(req):return httpx.Response(429,headers={'x-ratelimit-secret':secrets[1],'retry-after':'1'},json={})
    scheduler,pool=start(s,run,jobs,handler);scheduler.join(5)
    for j in jobs:
        for b in s.batches(j):assert not any(k in json.dumps(b['diagnostics']) for k in secrets)
    assert not any(k in json.dumps(s.log.snapshot()) for k in secrets)
    for key in s.kira_keys():s.change_kira_key(key['id'],delete=True)
    assert not s.kira_keys() and s.load_key() is None
    s.close()


def test_shared_rate_cooldown_survives_restart(tmp_path):
    from transcomic.providers import BatchResult
    s,_,_,_=setup(tmp_path,2,1,files=1)
    pool=KiraPoolProvider(s,'mock');key,adapter=pool.reserve(1,1)
    pool.release(key,BatchResult({},'Rate limited',diagnostics={'failure_kind':'rate_limit','retry_after_seconds':60}))
    s.close();s=Store(tmp_path/'state.db');pool=KiraPoolProvider(s,'mock')
    assert pool.reserve(1,1) is None
    s.close()


def test_100_file_queue_work_stealing_and_ready_no_keys(tmp_path):
    s,docs,run,jobs=setup(tmp_path,0,2,files=100,batches=1)
    calls=[]
    async def handler(req):
        calls.append(tuple(r['id'] for r in lines(req)));await asyncio.sleep(.002);return success(req)
    scheduler,pool=start(s,run,jobs,handler)
    scheduler.join(3)
    assert not scheduler._thread.is_alive() and not scheduler._paused.is_set()
    assert not calls
    s.add_kira_keys('sk-added-A sk-added-B sk-added-C')
    time.sleep(.05);assert not calls
    scheduler.start_many(jobs,run);scheduler.join(15)
    assert not scheduler._thread.is_alive() and len(calls)==100 and len(set(calls))==100
    assert s.run(run)['state']=='finished'
    s.close()


def test_disable_enable_cannot_clear_confirmed_quota(tmp_path):
    s=Store(tmp_path/'s.db');s.add_kira_keys('sk-safe-synthetic')
    key=s.kira_keys()[0]['id'];s.kira_key_status(key,'QUOTA_EXHAUSTED')
    s.change_kira_key(key,enabled=False);assert s.kira_keys()[0]['status']=='DISABLED'
    s.change_kira_key(key,enabled=True);assert s.kira_keys()[0]['status']=='QUOTA_EXHAUSTED'
    s.change_kira_key(key,restore=True);assert s.kira_keys()[0]['status']=='READY'
    s.close()
