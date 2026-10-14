import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from transcomic.documents import ImportedText, TextLine
from transcomic.docx_document import UnsafeDocx, import_docx, inspect_docx, render_docx
from transcomic.storage import Store


def _fixture(path: Path) -> bytes:
    image_path = path.with_suffix(".png")
    Image.new("RGB", (16, 12), "#cc8844").save(image_path)
    document = Document()
    p = document.add_paragraph()
    p.add_run("Hello ")
    p.add_run("thế giới").bold = True
    document.add_paragraph("")
    p = document.add_paragraph()
    p.add_run("First")
    p.add_run().add_break()
    p.add_run("Second")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Ô B"
    table.cell(1, 0).text = "Cell C"
    table.cell(1, 1).text = "Cell D"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    document.add_picture(str(image_path))
    document.save(path)
    return path.read_bytes()


def _store_docx(store: Store, path: Path, raw: bytes) -> str:
    parsed = inspect_docx(raw)
    data = ImportedText(tuple(TextLine(unit.text, "") for unit in parsed.units), "docx", raw)
    doc_id, created = store.import_document(path, data, document_format="docx", locators=[u.locator for u in parsed.units])
    assert created
    return doc_id


def test_docx_roundtrip_structure_and_snapshot(tmp_path: Path):
    path = tmp_path / "chapter.docx"
    raw = _fixture(path)
    parsed = import_docx(path)
    assert "Hello thế giới" in [u.text for u in parsed.units]
    assert "First" in [u.text for u in parsed.units]
    assert "Second" in [u.text for u in parsed.units]
    assert "Nested" in [u.text for u in parsed.units]
    store = Store(tmp_path / "app.db")
    doc_id = _store_docx(store, path, raw)
    path.unlink()
    store.close()
    reopened = Store(tmp_path / "app.db")
    rows = reopened.units(doc_id)
    assert len(rows) == len(parsed.units)
    first = next(row for row in rows if row["source"] == "Hello thế giới")
    reopened.edit_line(first["id"], "translation", "Xin chào thế giới")
    content, missing = render_docx(reopened.document(doc_id)["snapshot"], reopened.units(doc_id))
    assert missing > 0
    output = Document(io.BytesIO(content))
    assert output.paragraphs[0].text == "Xin chào thế giới"
    assert output.paragraphs[1].text == ""
    assert output.paragraphs[2].text == "First\nSecond"
    assert len(output.tables) == 1
    assert output.tables[0].cell(0, 0).text == "Cell A"
    assert output.tables[0].cell(1, 1).tables[0].cell(0, 0).text == "Nested"
    with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(content)) as after:
        media = [name for name in before.namelist() if name.startswith("word/media/")]
        assert media
        assert all(before.read(name) == after.read(name) for name in media)
        assert set(before.namelist()) == set(after.namelist())
    reopened.close()


def test_docx_rejects_unsupported_header(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Body")
    doc.sections[0].header.paragraphs[0].text = "Important header text"
    path = tmp_path / "header.docx"
    doc.save(path)
    with pytest.raises(UnsafeDocx, match="header"):
        import_docx(path)


def test_docx_rejects_field_and_invalid_file(tmp_path: Path):
    doc = Document()
    run = doc.add_paragraph("Hello").add_run()
    field = OxmlElement("w:fldChar")
    field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "begin")
    run._r.append(field)
    path = tmp_path / "field.docx"
    doc.save(path)
    with pytest.raises(UnsafeDocx, match="fldChar"):
        import_docx(path)
    with pytest.raises(UnsafeDocx):
        inspect_docx(b"not a docx")


def test_docx_rejects_missing_or_duplicate_mapping(tmp_path: Path):
    path = tmp_path / "chapter.docx"
    raw = _fixture(path)
    store = Store(tmp_path / "s.db")
    doc_id = _store_docx(store, path, raw)
    rows = store.units(doc_id)
    with pytest.raises(UnsafeDocx, match="khớp"):
        render_docx(raw, rows[:-1])
    rows[1]["locator"] = rows[0]["locator"]
    with pytest.raises(UnsafeDocx, match="trùng"):
        render_docx(raw, rows)
    store.close()


def test_docx_preserves_inline_image_and_incomplete_source(tmp_path: Path):
    image_path = tmp_path / "inline.png"
    Image.new("RGB", (12, 12), "#aa2244").save(image_path)
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Before")
    paragraph.add_run().add_picture(str(image_path))
    paragraph.add_run("After")
    path = tmp_path / "inline.docx"
    doc.save(path)
    raw = path.read_bytes()
    store = Store(tmp_path / "inline.db")
    doc_id = _store_docx(store, path, raw)
    rows = store.units(doc_id)
    assert [row["source"] for row in rows] == ["Before", "After"]
    store.edit_line(rows[0]["id"], "translation", "Trước")
    output, missing = render_docx(store.document(doc_id)["snapshot"], store.units(doc_id))
    assert missing == 1
    assert Document(io.BytesIO(output)).paragraphs[0].text == "TrướcAfter"
    with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
        media = next(name for name in before.namelist() if name.startswith("word/media/"))
        assert before.read(media) == after.read(media)
        xml = after.read("word/document.xml").decode()
        assert xml.index("Trước") < xml.index("<w:drawing") < xml.index("After")
    store.close()


def test_docx_rejects_tracked_change(tmp_path: Path):
    doc = Document()
    paragraph = doc.add_paragraph("Visible")
    inserted = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Tracked"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    path = tmp_path / "tracked.docx"
    doc.save(path)
    with pytest.raises(UnsafeDocx, match="ins"):
        import_docx(path)


def test_docx_blocks_invalid_xml_translation(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Hello")
    path = tmp_path / "invalid_text.docx"
    doc.save(path)
    store = Store(tmp_path / "invalid_text.db")
    doc_id = _store_docx(store, path, path.read_bytes())
    row = store.units(doc_id)[0]
    store.edit_line(row["id"], "translation", "Bad\x01text")
    with pytest.raises(UnsafeDocx, match="XML Word"):
        render_docx(store.document(doc_id)["snapshot"], store.units(doc_id))
    store.close()


def test_long_docx_keeps_paragraph_order(tmp_path: Path):
    doc = Document()
    for number in range(500):
        doc.add_paragraph(f"Đoạn {number}")
    path = tmp_path / "long.docx"
    doc.save(path)
    store = Store(tmp_path / "long.db")
    doc_id = _store_docx(store, path, path.read_bytes())
    rows = store.units(doc_id)
    assert len(rows) == 500
    store.edit_line(rows[250]["id"], "translation", "Đoạn giữa đã dịch")
    output, missing = render_docx(store.document(doc_id)["snapshot"], store.units(doc_id))
    paragraphs = Document(io.BytesIO(output)).paragraphs
    assert len(paragraphs) == 500 and missing == 499
    assert paragraphs[0].text == "Đoạn 0"
    assert paragraphs[250].text == "Đoạn giữa đã dịch"
    assert paragraphs[-1].text == "Đoạn 499"
    store.close()
