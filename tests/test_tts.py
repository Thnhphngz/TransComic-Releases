"""Offline TTS safety tests use synthetic audio, never user documents or APIs."""
import json
from pathlib import Path
import threading
import time
import pytest
from transcomic.tts.documents import segments,scan,read_document
from transcomic.tts.engine import chunks
from transcomic.tts.storage import TtsStore
from transcomic.tts.output import Outputs,base_name,MANIFEST


def make(tmp_path,text='One\nTwo',name='Chap 9'):
    store=TtsStore(tmp_path/'state.sqlite3');run=store.draft()
    store.add_document(run,dict(name=name,text=text,path=None,line_count=len(text.splitlines())))
    return store,run


def prepare(tmp_path,text='One\nTwo',name='Chap 9',format='wav'):
    store,run=make(tmp_path,text,name);root=tmp_path/'output';root.mkdir()
    config=dict(device='CPU',voice='af_heart',speed=1,language='en-us',variant='standard',format=format)
    output=Outputs(store);plan=output.plan(run,'custom',str(root),format);output.confirm(run,config,plan)
    return store,run,output,root


def wave():
    np=pytest.importorskip('numpy');pytest.importorskip('soundfile')
    return np.sin(np.arange(2400)*.07).astype('float32')*.2


class FakeEngine:
    barrier=None;calls=[]
    def __init__(self,manager,variant,backend,device_id):self.backend=backend
    def probe(self,*args):return 1
    def synthesize(self,text,*args):
        self.calls.append((self.backend,text))
        if self.barrier:self.barrier.wait(timeout=5)
        time.sleep(.002 if self.backend=='CPUExecutionProvider' else .005)
        yield wave()


def finish(store,run,output):
    from transcomic.tts.audio import encode
    while item:=store.claim(run):
        folder=Path(store.doc(item['document_id'])['folder']);temp=folder/(item['attempt']+'.part')
        info=encode(temp,[wave()]);assert output.commit(item,temp,info,'mock-CPU')
    output.finalize(run)


@pytest.mark.parametrize('mode,expected',[
    ('lines',['Mr. Smith 3.14... Why!?','Hello. End']),
    ('sentences',['Mr.','Smith 3.','14...','Why!?','Hello.','End\n.']),
    ('combined',['Mr.','Smith 3.','14...','Why!?','Hello.','End'])])
def test_literal_split(mode,expected):
    text='\nMr. Smith 3.14... Why!?\n\nHello. End\n. ! ?'
    items=segments(text,mode);assert [i['text'] for i in items]==expected
    for i in items:assert text[i['start']:i['end']]==i['text']
    assert items[0]['line_start']==2


def test_long_chunks_and_10000_source_positions():
    text=('Word . 3.14\n'*1000)+'最後'
    assert ''.join(chunks(text))==text
    items=segments('A\r\n'*10000+'Last','combined')
    assert items[-1]['line_start']==10001 and items[-1]['text']=='Last'
    assert segments(' . ! ? \n')==[]


def test_import_and_scan_docx_unicode_manual_break(tmp_path):
    from docx import Document
    a=tmp_path/'Chương.txt';a.write_text('First\n\nTiếng Việt',encoding='utf8')
    doc=Document();p=doc.add_paragraph('First');p.add_run().add_break();p.add_run('Second');doc.add_paragraph('最後')
    b=tmp_path/'chapter.docx';doc.save(b)
    before=b.read_bytes();txt=read_document(a);word=read_document(b)
    assert txt['line_count']==3 and '\n\n' in txt['text']
    assert 'First\nSecond' in word['text'] and len(word['locators'])==3
    files,errors=scan([tmp_path,a,tmp_path/'missing']);assert len(files)==2 and len(errors)==1 and b.read_bytes()==before


def test_draft_edits_duplicates_skip_merge_split_persistence(tmp_path):
    store,run=make(tmp_path,'Hello world\nSecond sentence')
    first=store.items(run)[0];store.manual_split(first['id'],6)
    items=store.items(run);assert [i['text'].strip() for i in items]==['Hello','world','Second sentence']
    store.merge([items[0]['id'],items[1]['id']]);store.edit(first['id'],'Changed')
    assert store.items(run)[0]['revision']>=2
    with pytest.raises(ValueError):store.manual_split(first['id'],0)
    store.close();store=TtsStore(tmp_path/'state.sqlite3');assert store.items(run)[0]['text']=='Changed';store.close()


@pytest.mark.parametrize('value',['CON','nul.txt','COM1','bad/name','a.','../bad','', 'a'*101])
def test_reject_unsafe_folder_names(value):
    with pytest.raises(ValueError):base_name(value)


@pytest.mark.parametrize('count',[1,120,1001])
def test_exact_folder_count_and_resume_skip_valid(tmp_path,count):
    store,run,output,root=prepare(tmp_path,'\n'.join(f'Line {i}' for i in range(count)))
    finish(store,run,output);folder=root/f'Chap 9_{count}tts'
    assert folder.is_dir() and len(list(folder.glob('*.wav')))==count
    assert (folder/'0001.wav').is_file();assert store.items(run)[-1]['filename']==f'{count:04}.wav'
    output.reconcile(run);assert store.claim(run) is None
    assert len(json.loads((folder/MANIFEST).read_text(encoding='utf8'))['segments'])==count
    store.close()


def test_plan_zero_duplicate_names_unicode_and_collisions(tmp_path):
    store,run=make(tmp_path,' . ! ?')
    with pytest.raises(ValueError):Outputs(store).plan(run,'custom',str(tmp_path),'wav')
    store.add_document(run,dict(name='Chương',text='Yes',path=None,line_count=1))
    store.add_document(run,dict(name='Chương',text='No',path=None,line_count=1))
    plan=Outputs(store).plan(run,'custom',str(tmp_path),'wav');assert len(plan)==2 and plan[1]['base']=='Chương_2'
    Path(plan[0]['folder']).mkdir()
    revised=Outputs(store).plan(run,'custom',str(tmp_path),'wav')
    assert revised[0]['folder']!=plan[0]['folder'] and Path(plan[0]['folder']).is_dir()
    with pytest.raises(ValueError):Outputs(store).plan(run,'beside',None,'wav')
    with pytest.raises(ValueError):Outputs(store).plan(run,'custom','','wav')
    store.close()


def test_stale_worker_cannot_overwrite_and_confirmed_regeneration(tmp_path):
    from transcomic.tts.audio import encode
    store,run,output,root=prepare(tmp_path,'Original');item=store.claim(run);folder=Path(store.docs(run)[0]['folder'])
    temp=folder/'stale.part';info=encode(temp,[wave()]);store.edit(item['id'],'Changed')
    assert not output.commit(item,temp,info,'mock') and not (folder/'0001.wav').exists()
    finish(store,run,output);folder=root/'Chap 9_1tts';before=(folder/'0001.wav').read_bytes()
    id=store.items(run)[0]['id'];store.edit(id,'Regenerate');item=store.claim(run);temp=folder/'new.part';info=encode(temp,[wave()*.5])
    with pytest.raises(FileExistsError):output.commit(item,temp,info,'mock')
    assert (folder/'0001.wav').read_bytes()==before
    output.retry(run,[id]);item=store.claim(run);folder=Path(store.docs(run)[0]['folder']);temp=folder/'new2.part';info=encode(temp,[wave()*.5]);assert output.commit(item,temp,info,'mock')
    output.finalize(run);assert (root/'Chap 9_1tts'/'0001.wav').read_bytes()!=before;store.close()


def test_crash_after_replace_before_db_commit_and_relink(tmp_path,monkeypatch):
    from transcomic.tts.audio import encode
    store,run,output,root=prepare(tmp_path,'One');item=store.claim(run);folder=Path(store.docs(run)[0]['folder']);temp=folder/'pending.part';info=encode(temp,[wave()])
    original=store.execute
    def crash(sql,args=()):
        if "state='success'" in sql:raise RuntimeError('simulated power loss')
        return original(sql,args)
    monkeypatch.setattr(store,'execute',crash)
    with pytest.raises(RuntimeError):output.commit(item,temp,info,'mock')
    monkeypatch.setattr(store,'execute',original);store.close();store=TtsStore(tmp_path/'state.sqlite3');output=Outputs(store)
    assert store.items(run)[0]['state']=='unknown';output.reconcile(run);assert store.items(run)[0]['state']=='success'
    output.finalize(run);doc=store.docs(run)[0];old=Path(doc['folder']);moved=root/'Moved';old.rename(moved)
    with pytest.raises(FileNotFoundError):output.reconcile(run)
    output.relink(doc['id'],moved);assert store.docs(run)[0]['folder']==str(moved)
    store.delete_history([run]);assert (moved/'0001.wav').exists();store.close()


def test_usb_loss_corrupt_audio_and_no_silent_retry(tmp_path):
    store,run,output,root=prepare(tmp_path,'One');finish(store,run,output);folder=Path(store.docs(run)[0]['folder'])
    (folder/'0001.wav').write_bytes(b'broken');output.reconcile(run)
    assert store.items(run)[0]['state']=='error' and store.claim(run) is None
    folder.rename(root/'unplugged')
    with pytest.raises(FileNotFoundError):output.reconcile(run)
    assert not folder.exists();store.close()


def test_hybrid_dynamic_claims_no_duplicates_and_pause(tmp_path,monkeypatch):
    from transcomic.tts.scheduler import Scheduler
    import transcomic.tts.scheduler as module
    store,run,output,root=prepare(tmp_path,'\n'.join(f'Line {i}' for i in range(121)))
    monkeypatch.setattr(module,'backend_plan',lambda _:['CPUExecutionProvider','mock-GPU'])
    FakeEngine.calls=[];scheduler=Scheduler(store,None,engine_factory=FakeEngine);scheduler.start(run);scheduler.thread.join(40)
    assert not scheduler.busy and store.run(run)['state']=='completed'
    assert len(FakeEngine.calls)==121 and len({t for _,t in FakeEngine.calls})==121
    assert {b for b,_ in FakeEngine.calls}=={'CPUExecutionProvider','mock-GPU'}
    assert len(list((root/'Chap 9_121tts').glob('*.wav')))==121;store.close()


def test_pause_stop_resume_and_no_db_lock_during_inference(tmp_path):
    from transcomic.tts.scheduler import Scheduler
    entered=threading.Event();release=threading.Event()
    class Slow(FakeEngine):
        def synthesize(self,*args):entered.set();assert release.wait(5);yield wave()
    store,run,output,root=prepare(tmp_path,'One\nTwo');scheduler=Scheduler(store,None,engine_factory=Slow)
    scheduler.start(run);assert entered.wait(5)
    # A second SQLite writer works while model inference is waiting.
    import sqlite3
    connection=sqlite3.connect(store.path,timeout=.2);connection.execute("INSERT INTO preferences VALUES ('check','ok')");connection.commit();connection.close()
    scheduler.pause();release.set();scheduler.thread.join(10)
    assert store.run(run)['state']=='paused' and store.progress(run)=={'pending':1,'success':1}
    scheduler.start(run);scheduler.thread.join(10);assert store.run(run)['state']=='completed';store.close()


@pytest.mark.parametrize('format',['wav','mp3'])
def test_real_encoder_decode_and_silence(format,tmp_path):
    from transcomic.tts.audio import encode,validate
    if format=='mp3':pytest.importorskip('lameenc')
    path=tmp_path/('encoded.'+format);info=encode(path,[wave(),wave()],format)
    assert .15<info['duration']<.4 and validate(path)['sha256']==info['sha256']
    if format=='mp3':assert path.read_bytes()[:4]!=b'RIFF'
    with pytest.raises(ValueError):encode(tmp_path/'silent.part',[wave()*0],format)


def test_auto_fallback_and_explicit_gpu_failure_preserves_progress(tmp_path,monkeypatch):
    from transcomic.tts.scheduler import Scheduler
    import transcomic.tts.scheduler as module
    class ProbeFailure(FakeEngine):
        def probe(self,*args):
            if self.backend!='CPUExecutionProvider':raise RuntimeError('VRAM / unsupported graph')
            return 1
    store,run,output,root=prepare(tmp_path,'One');config=json.loads(store.run(run)['config']);config['device']='Auto';store.execute('UPDATE runs SET config=? WHERE id=?',(json.dumps(config),run))
    monkeypatch.setattr(module,'backend_plan',lambda _:['mock-GPU']);scheduler=Scheduler(store,None,engine_factory=ProbeFailure)
    scheduler.start(run);scheduler.thread.join(10);assert store.run(run)['state']=='completed' and 'CPU' in scheduler.message
    assert store.items(run)[0]['backend']=='CPUExecutionProvider';store.close()


def test_same_revision_explicit_retry_really_regenerates(tmp_path):
    store,run,output,root=prepare(tmp_path,'One');finish(store,run,output);item=store.items(run)[0]
    output.retry(run,[item['id']]);output.reconcile(run)
    claimed=store.claim(run);assert claimed and claimed['id']==item['id'] and claimed['overwrite_allowed']
    store.close()


@pytest.mark.parametrize('reason',[PermissionError('permission'),OSError(28,'disk full'),FileNotFoundError('USB unplugged')])
def test_write_failure_pauses_no_output_fallback(tmp_path,monkeypatch,reason):
    from transcomic.tts.scheduler import Scheduler
    import transcomic.tts.scheduler as module
    store,run,output,root=prepare(tmp_path,'One\nTwo')
    def fail(*args,**kwargs):raise reason
    monkeypatch.setattr(module,'encode',fail);scheduler=Scheduler(store,None,engine_factory=FakeEngine)
    scheduler.start(run);scheduler.thread.join(10)
    assert store.run(run)['state']=='paused' and store.progress(run)=={'error':1,'pending':1}
    assert not list(root.rglob('*.wav')) and store.docs(run)[0]['folder'].startswith(str(root));store.close()


def test_model_download_cancel_corrupt_hash_cached_and_remove_safe(tmp_path,monkeypatch):
    import hashlib,httpx
    import transcomic.tts.models as module
    payload=b'verified model fixture';name='onnx/model.onnx';cancel=threading.Event();requests=[]
    monkeypatch.setattr(module,'ASSETS',{name:dict(size=len(payload),sha256=hashlib.sha256(payload).hexdigest())})
    manager=module.ModelManager(tmp_path/'Models');monkeypatch.setattr(manager,'required',lambda _: [name])
    def handler(request):requests.append(str(request.url));return httpx.Response(200,content=payload)
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        cancel.set()
        with pytest.raises(module.Cancelled):manager.install(cancel=cancel,client=client)
        cancel.clear();manager.install(client=client);assert manager.valid(name)
        manager.install(client=client);assert len(requests)==1 and module.REVISION in requests[0]
    unrelated=manager.root/'user-audio.wav';unrelated.write_bytes(b'keep')
    manager.remove();assert unrelated.read_bytes()==b'keep' and not manager.path(name).exists()
    with httpx.Client(transport=httpx.MockTransport(lambda _:httpx.Response(200,content=b'x'*len(payload)))) as client:
        with pytest.raises(ValueError):manager.install(client=client)
    assert not manager.valid(name) and not manager.path(name).exists()


def test_final_folder_collision_keeps_valid_audio_and_no_success_claim(tmp_path):
    store,run,output,root=prepare(tmp_path,'One');(root/'Chap 9_1tts').mkdir()
    finish(store,run,output)
    assert (root/'Chap 9_2_1tts'/'0001.wav').is_file() and store.docs(run)[0]['folder']==str(root/'Chap 9_2_1tts')
    assert list((root/'Chap 9_1tts').iterdir())==[];store.close()


def test_symlink_or_junction_rejected(tmp_path):
    from transcomic.tts.output import safe_path
    target=tmp_path/'real';target.mkdir();link=tmp_path/'link'
    try:link.symlink_to(target,target_is_directory=True)
    except OSError:pytest.skip('Windows account has no symlink privilege')
    with pytest.raises(ValueError):safe_path(link/'output')


def test_packaging_audit_never_builds():
    from scripts.tts_packaging import audit
    result=audit();assert isinstance(result['missing'],list) and 'license_gate' in result


def test_explicit_gpu_failure_never_leases_or_claims_gpu_success(tmp_path,monkeypatch):
    from transcomic.tts.scheduler import Scheduler
    import transcomic.tts.scheduler as module
    class Broken(FakeEngine):
        def probe(self,*args):raise RuntimeError('unsupported GPU / VRAM')
    store,run,output,root=prepare(tmp_path,'One\nTwo');config=json.loads(store.run(run)['config']);config['device']='GPU'
    store.execute('UPDATE runs SET config=? WHERE id=?',(json.dumps(config),run));monkeypatch.setattr(module,'backend_plan',lambda _:['mock-GPU'])
    scheduler=Scheduler(store,None,engine_factory=Broken);scheduler.start(run);scheduler.thread.join(5)
    assert store.run(run)['state']=='paused' and store.progress(run)=={'pending':2}
    assert not list(root.rglob('*.wav'));store.close()


@pytest.mark.parametrize('profile_provider',['CPUExecutionProvider','CUDAExecutionProvider'])
def test_gpu_requires_node_proof_cuda_index_and_no_hidden_fallback(tmp_path,monkeypatch,profile_provider):
    from types import SimpleNamespace
    import sys
    np=pytest.importorskip('numpy');captured={}
    class Session:
        def __init__(self,path,sess_options,providers):self.options=sess_options;captured['providers']=providers
        def get_providers(self):return ['CUDAExecutionProvider','CPUExecutionProvider']
        def disable_fallback(self):captured['disabled_fallback']=True
        def end_profiling(self):
            path=Path(self.options.profile_file_prefix+'.json');path.write_text(json.dumps([{'args':{'provider':profile_provider}}]));return str(path)
    class Kokoro:
        @staticmethod
        def from_session(session,*args):return SimpleNamespace(sess=session,_input_dtypes={'speed':np.dtype('float32')},create=lambda *args,**kwargs:(wave(),24000))
    rt=SimpleNamespace(SessionOptions=SimpleNamespace,InferenceSession=Session,get_available_providers=lambda:['CUDAExecutionProvider','CPUExecutionProvider'])
    monkeypatch.setitem(sys.modules,'onnxruntime',rt);monkeypatch.setitem(sys.modules,'kokoro_onnx',SimpleNamespace(Kokoro=Kokoro))
    monkeypatch.setattr('transcomic.tts.adapters.preferred_device',lambda:1)
    manager=SimpleNamespace(ready=lambda _:True,path=lambda _:tmp_path/'model.onnx',voices_pack=lambda:tmp_path/'voices.npz')
    from transcomic.tts.engine import KokoroEngine
    engine=KokoroEngine(manager,backend='CUDAExecutionProvider');options=captured['providers'][0][1]
    assert options['device_id']==0 and options['gpu_mem_limit']==2*1024**3 and captured['disabled_fallback']
    if profile_provider=='CPUExecutionProvider':
        with pytest.raises(RuntimeError,match='GPU'):engine.probe('af_heart',1,'en-us')
    else:assert engine.probe('af_heart',1,'en-us')>0 and engine.gpu_verified
