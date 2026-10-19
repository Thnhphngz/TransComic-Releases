"""Integration fixtures: real services, mock HTTP and tiny synthetic PCM only."""
import asyncio
from collections import Counter
import json
from pathlib import Path
import threading
import time
from types import SimpleNamespace
import httpx
import numpy as np
import pytest
from PySide6.QtGui import QImage,QColor
from transcomic.storage import Store
from transcomic.automation.repository import Repository,file_hash
from transcomic.automation.validation import defaults
from transcomic.automation.adapters import Translation,Kokoro,Capcut
from transcomic.automation.controller import Controller
from transcomic.kira_pool import KiraPoolProvider
from transcomic.tts.storage import TtsStore
from test_kira_auto_recovery import response,request_lines


class FakeEngine:
    calls=[];fail=None;entered=None;release=None
    live=0;max_live=0
    def __init__(self,*args):
        self.backend='CPUExecutionProvider';type(self).live+=1;type(self).max_live=max(type(self).max_live,type(self).live)
    def probe(self,*args):pass
    def synthesize(self,text,*args):
        self.calls.append(text)
        if self.entered:self.entered.set();self.release.wait(4)
        if type(self).fail and type(self).fail(text):raise ValueError('Synthetic local WAV failure')
        yield np.sin(np.arange(2400,dtype=np.float32)*.12)*.1
    def close(self):type(self).live-=1


@pytest.fixture
def setup(tmp_path):
    store=Store(tmp_path/'state.sqlite3');repo=Repository(store);tts=TtsStore(tmp_path/'TTS/tts.sqlite3');repo.tts_store=tts;tts.automation=repo
    output=tmp_path/'capcut';output.mkdir();cfg=defaults(store);cfg.update(model='mock',output=str(output),device='CPU')
    id=repo.create('Synthetic story',cfg);store.add_kira_keys('automation-synthetic-key')
    FakeEngine.calls=[];FakeEngine.fail=None;FakeEngine.entered=None;FakeEngine.release=None;FakeEngine.live=0;FakeEngine.max_live=0
    def provider(store,job):return KiraPoolProvider(store,'mock',transport=httpx.MockTransport(response),rpm=100000,tpm=100000000)
    services=SimpleNamespace(tts_store=tts,translation=Translation(repo,provider,network_probe=lambda provider:True),kokoro=Kokoro(repo,tts,None,engine_factory=FakeEngine),capcut=Capcut(repo,tts),preflight=lambda cfg:None)
    controller=Controller(repo,services)
    services.translation.backoff=lambda phase:0;services.kokoro.backoff=lambda phase:0
    yield repo,id,services,controller
    controller.stop_event.set();services.translation.stop();controller.join(6)
    assert not controller.busy
    from transcomic.tts.runtime import logger
    log=logger(tts.path.parent/'Logs')
    for handler in list(log.handlers):handler.close();log.removeHandler(handler)
    tts.close();store.close()


def add_group(repo,id,root,name='1-2',count=3,images=None):
    root.mkdir(parents=True,exist_ok=True);path=root/(name+'.txt');path.write_text('\n'.join(f'Văn bản tổng hợp số {i}' for i in range(count)),encoding='utf8')
    gid=repo.add(id,[path])[0];folder=root/(name+'-images');folder.mkdir()
    for n in range(count if images is None else images):
        image=QImage(20,30,QImage.Format_RGB32);image.fill(QColor(20+n,40,80));assert image.save(str(folder/f'{n+1:04}.png'))
    repo.pair_images(gid,folder);return gid,path,folder


def run(controller,id,timeout=10):
    proof=controller.validate(id);assert proof['passed'],proof['errors']+[g['errors'] for g in proof['groups']]
    controller.start(id);controller.join(timeout);assert not controller.busy


def test_import_duplicate_same_name_distinct_paths_natural_order_and_metadata_delete(setup,tmp_path):
    repo,id,s,c=setup
    a,path,_=add_group(repo,id,tmp_path/'one','11-12');b,path2,_=add_group(repo,id,tmp_path/'two','1-2')
    assert repo.add(id,[path,path2])==[]
    other,otherpath,_=add_group(repo,id,tmp_path/'three','1-2')
    assert len(repo.groups(id))==3 and len({g['document_id'] for g in repo.groups(id)})==3
    repo.move(other,-1);assert repo.groups(id)[0]['id']==other
    with pytest.raises(ValueError):repo.add(id,[tmp_path/'one'])
    before=[file_hash(p) for p in (path,path2,otherpath)]
    repo.delete(id)
    assert [file_hash(p) for p in (path,path2,otherpath)]==before
    assert len(repo.store.documents())==3


@pytest.mark.parametrize('images',[0,2,4])
def test_first_start_requires_all_groups_valid(setup,tmp_path,images):
    repo,id,s,c=setup;add_group(repo,id,tmp_path/'a',images=images)
    proof=c.validate(id);assert not proof['passed']
    with pytest.raises(ValueError):c.start(id)
    assert not repo.workset(id)['started'] and FakeEngine.calls==[]


@pytest.mark.parametrize('mode,count',[('separate',3),('merged',1),('split',2)])
def test_real_pipeline_all_modes_stable_mapping_no_source_changes(setup,tmp_path,mode,count):
    repo,id,s,c=setup;cfg=repo.workset(id)['config'];cfg.update(capcut_mode=mode,part_count=2,project_name='Synthetic');repo.configure(id,cfg)
    sources=[]
    for i,n in enumerate((3,4,2)):
        gid,path,images=add_group(repo,id,tmp_path/f'g{i}',f'{i+1}-{i+2}',n);sources.append((path,file_hash(path)))
    run(c,id)
    assert repo.workset(id)['state']=='completed',[(e['stage'],e['message']) for e in repo.events(id)]
    projects=[a for a in repo.artifacts(id) if a['kind']=='capcut'];assert len(projects)==count and all(a['state']=='complete' for a in projects)
    part_order={p['id']:i for i,p in enumerate(repo.workset(id)['checkpoint']['parts'])}
    projects.sort(key=lambda a:part_order[a['unit']])  # Independent parts may finish in any order.
    actual=[]
    for a in projects:
        path=Path(a['path']);mapping=json.loads((path/'transcomic_mapping.json').read_text(encoding='utf8'))
        actual.extend((r['group_id'],int(Path(r['audio']).stem)) for r in mapping['pairs'])
        assert (path/'draft_cover.jpg').is_file()
        assert not list(path.rglob('*.wav')) and not list(path.rglob('*.png'))
    expected=[(g['id'],n) for g in repo.groups(id) for n in range(1,len(g['data']['slots'])+1)]
    assert actual==expected and len(FakeEngine.calls)==9
    assert all(file_hash(p)==h for p,h in sources) and FakeEngine.max_live==1
    assert len(list(tmp_path.rglob('*.en.txt')))==3
    assert len(list(tmp_path.rglob('TransComic_Automation_Reports/*.json')))==1


def test_resume_missing_wav_and_manual_editor_only_invalidates_one_voice(setup,tmp_path):
    repo,id,s,c=setup;gid,path,images=add_group(repo,id,tmp_path/'a');run(c,id)
    group=repo.group(gid);runid=group['data']['tts_run'];items=s.tts_store.items(runid);folder=Path(s.tts_store.docs(runid)[0]['folder'])
    old={p:file_hash(p) for p in folder.glob('*.wav')};first_project=Path(repo.artifacts(id)[-1]['path']);old_projects={a['path']:a['sha256'] for a in repo.artifacts(id) if a['kind']=='capcut'}
    (folder/items[1]['filename']).unlink();FakeEngine.calls.clear();run(c,id)
    assert len(FakeEngine.calls)==1 and repo.workset(id)['state']=='completed'
    line=repo.store.lines(group['document_id'])[0];repo.store.edit_line(line['id'],'translation','The editor has changed this sentence')
    assert repo.group(gid)['data']['state']=='stale' and s.tts_store.items(runid)[0]['state']=='stale'
    calls_before=sum(len(repo.store.attempts(b['id'])) for j in repo.store.jobs() for b in repo.store.batches(j['id']))
    FakeEngine.calls.clear();run(c,id)
    assert FakeEngine.calls==['The editor has changed this sentence']
    assert sum(len(repo.store.attempts(b['id'])) for j in repo.store.jobs() for b in repo.store.batches(j['id']))==calls_before
    assert repo.store.line(line['id'])['translation']=='The editor has changed this sentence'
    assert file_hash(folder/items[0]['filename'])==old[folder/items[0]['filename']]
    assert len([a for a in repo.artifacts(id) if a['kind']=='capcut'])>=2
    cfg=repo.workset(id)['config'];cfg['speed']=1.2
    with pytest.raises(ValueError):repo.configure(id,cfg)


def test_safe_stop_fences_new_wavs_locks_edits_and_other_worksets(setup,tmp_path):
    repo,id,s,c=setup;gid,path,_=add_group(repo,id,tmp_path/'a',count=4)
    FakeEngine.entered=threading.Event();FakeEngine.release=threading.Event()
    assert c.validate(id)['passed'];c.start(id);assert FakeEngine.entered.wait(4)
    second=repo.create('Second',repo.workset(id)['config'])
    with pytest.raises(ValueError):repo.activate(second)
    line=repo.store.lines(repo.group(gid)['document_id'])[0]
    with pytest.raises(ValueError):repo.store.edit_line(line['id'],'translation','Forbidden while running')
    c.stop();assert repo.workset(id)['state']=='stopping'
    with pytest.raises(ValueError):repo.configure(id,repo.workset(id)['config'])
    FakeEngine.release.set();c.join(6)
    assert not c.busy and repo.workset(id)['state']=='stopped'
    assert len(FakeEngine.calls)==1
    assert sum(i['state']=='success' for i in s.tts_store.items(repo.group(gid)['data']['tts_run']))==1
    assert repo.workset(second)['state']=='draft'


def test_wave_failures_keep_positions_and_max_five_attempts(setup,tmp_path,monkeypatch):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=3)
    # Engine returns the same sentence in this fixture; force the second invocation of each pass to fail by item ordinal.
    original=s.tts_store.claim
    def claim(*args,**kwargs):
        item=original(*args,**kwargs)
        FakeEngine.fail=lambda text:bool(item and item['position']==2)
        return item
    monkeypatch.setattr(s.tts_store,'claim',claim)
    run(c,id,timeout=22)
    assert repo.workset(id)['state']=='needs_attention'
    meta=repo.group(gid)['data'];items=s.tts_store.items(meta['tts_run']);folder=Path(s.tts_store.docs(meta['tts_run'])[0]['folder'])
    assert [i['state'] for i in items]==['success','error','success']
    assert (folder/'0001.wav').exists() and not (folder/'0002.wav').exists() and (folder/'0003.wav').exists()
    assert max(meta['wave_attempts'].values())==5 and len(FakeEngine.calls)==7
    assert not [a for a in repo.artifacts(id) if a['kind']=='capcut']


def test_source_manual_conflict_blocks_only_group_on_resume(setup,tmp_path):
    repo,id,s,c=setup;a,path,_=add_group(repo,id,tmp_path/'a');b,path2,_=add_group(repo,id,tmp_path/'b');run(c,id)
    line=repo.store.lines(repo.group(a)['document_id'])[0];repo.store.edit_line(line['id'],'translation','My manual sentence is protected')
    path.write_text('Nguồn khác\nVăn bản tổng hợp số 1\nVăn bản tổng hợp số 2',encoding='utf8')
    proof=c.validate(id);assert not proof['passed'] and not proof['errors']
    c.start(id);c.join(8)
    assert repo.group(a)['data']['state']=='needs_attention'
    assert repo.group(b)['data']['state']=='completed'
    assert repo.store.line(line['id'])['translation']=='My manual sentence is protected'


@pytest.mark.parametrize('failure',[504,'timeout',429,'schema'])
def test_translation_original_batches_cap_at_five_requests_and_no_provider_switch(setup,tmp_path,failure):
    from transcomic.providers import KiraProvider
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=6);calls=[]
    def handler(req):
        calls.append((json.loads(req.content)['model'],[r['id'] for r in request_lines(req)],time.monotonic()))
        if failure=='timeout':raise httpx.ReadTimeout('Synthetic request timeout',request=req)
        if failure=='schema':return response(req,[])
        return httpx.Response(failure,headers={'Retry-After':'0.03'} if failure==429 else {})
    s.translation.provider_factory=lambda store,job:KiraProvider('fixture','mock',transport=httpx.MockTransport(handler))
    run(c,id,timeout=12)
    assert repo.workset(id)['state']=='needs_attention' and FakeEngine.calls==[]
    assert len(calls)==10 and {model for model,ids,stamp in calls}=={'mock'}
    batches=repo.group(gid)['data']['batch_jobs'];assert len(batches)==2
    for jobs in batches.values():assert sum(len(repo.store.attempts(bid)) for bid in repo.batch_ids(jobs))==5
    assert all(set(ids)<=set(calls[0][1]) or set(ids)<=set(calls[1][1]) for model,ids,stamp in calls)
    if failure==429:
        per_batch=[stamp for model,ids,stamp in calls if ids==calls[0][1]]
        assert all(b-a>=.025 for a,b in zip(per_batch,per_batch[1:]))


def test_multi_group_translation_concurrency_is_one_global_ceiling(setup,tmp_path):
    repo,id,s,c=setup;add_group(repo,id,tmp_path/'a','1-2',61);add_group(repo,id,tmp_path/'b','3-4',61)
    active=0;maximum=0;calls=[]
    async def handler(req):
        nonlocal active,maximum
        active+=1;maximum=max(maximum,active);calls.append(request_lines(req))
        await asyncio.sleep(.04);active-=1;return response(req)
    s.translation.provider_factory=lambda store,job:KiraPoolProvider(store,'mock',transport=httpx.MockTransport(handler),rpm=100000,tpm=100000000)
    run(c,id,timeout=20)
    assert 1<maximum<=30 and len(calls)==42
    assert len(FakeEngine.calls)==122 and repo.workset(id)['state']=='completed'


def test_pipeline_overlaps_translation_and_tts_first_completed_group_first(setup,tmp_path):
    repo,id,s,c=setup;a,_,_=add_group(repo,id,tmp_path/'a','1-2',1);b,_,_=add_group(repo,id,tmp_path/'b','3-4',1)
    a_ids={r['id'] for r in repo.store.lines(repo.group(a)['document_id'])};seen_tts=threading.Event();times={}
    original=FakeEngine.synthesize
    def synthesize(engine,text,*args):
        times.setdefault('tts',time.monotonic());seen_tts.set();yield from original(engine,text,*args)
    FakeEngine.synthesize=synthesize
    async def handler(req):
        ids={r['id'] for r in request_lines(req)}
        if ids&a_ids:
            deadline=time.monotonic()+4
            while not seen_tts.is_set() and time.monotonic()<deadline:await asyncio.sleep(.01)
            times['later_translation']=time.monotonic()
        return response(req)
    s.translation.provider_factory=lambda store,job:KiraPoolProvider(store,'mock',transport=httpx.MockTransport(handler),rpm=100000,tpm=100000000)
    try:
        run(c,id)
        assert times['tts']<times['later_translation']
        ready=[e['group_id'] for e in reversed(repo.events(id)) if e['message'].startswith('Dịch hợp lệ')]
        assert ready[0]==b
    finally:FakeEngine.synthesize=original


def test_gemini_pipeline_uses_existing_provider_and_never_kira(setup,tmp_path):
    from transcomic.gemini import GeminiProvider,QuotaLimiter
    repo,id,s,c=setup;add_group(repo,id,tmp_path/'a');cfg=repo.workset(id)['config'];cfg.update(provider='gemini',model='gemini-2.0-flash');repo.configure(id,cfg)
    calls=[]
    def handler(req):
        calls.append(str(req.url));payload=json.loads(json.loads(req.content)['contents'][0]['parts'][0]['text']);rows=payload['lines']
        return httpx.Response(200,json={'candidates':[{'content':{'parts':[{'text':json.dumps({'lines':[{'id':r['id'],'text':'The translated story continues'} for r in rows]})}]},'finishReason':'STOP'}]})
    s.translation.provider_factory=lambda store,job:GeminiProvider([dict(id='fixture',secret='synthetic',project='fixture')],job['model'],limiter=QuotaLimiter(1000,100000000,1000),transport=httpx.MockTransport(handler))
    run(c,id)
    assert repo.workset(id)['state']=='completed' and len(calls)==1
    assert all('generativelanguage.googleapis.com' in url for url in calls)


def test_completed_unchanged_resume_is_noop_and_changed_image_only_rebuilds_project(setup,tmp_path):
    repo,id,s,c=setup;gid,path,folder=add_group(repo,id,tmp_path/'a');run(c,id)
    old_calls=list(FakeEngine.calls);artifacts=list(repo.artifacts(id));proof=c.validate(id)
    assert proof['passed'] and proof['pending'] is False
    with pytest.raises(ValueError,match='không có công việc'):c.start(id)
    assert repo.workset(id)['cycle']==1 and repo.artifacts(id)==artifacts
    image=QImage(25,30,QImage.Format_RGB32);image.fill(QColor('red'));assert image.save(str(folder/'0001.png'))
    run(c,id)
    assert FakeEngine.calls==old_calls
    assert len([a for a in repo.artifacts(id) if a['kind']=='capcut'])==2


def test_corrupt_wave_and_external_project_edits_are_preserved(setup,tmp_path):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a');run(c,id)
    meta=repo.group(gid)['data'];runid=meta['tts_run'];doc=s.tts_store.docs(runid)[0];wave=Path(doc['folder'])/'0002.wav';wave.write_bytes(b'external corrupt file')
    project=Path(meta['project']);(project/'external-edit.txt').write_text('Keep this user edit')
    before={p:file_hash(p) for p in project.rglob('*') if p.is_file()};FakeEngine.calls.clear();run(c,id)
    assert wave.read_bytes()==b'external corrupt file' and len(FakeEngine.calls)==1
    assert all(file_hash(p)==h for p,h in before.items())
    assert repo.group(gid)['data']['project']!=str(project)
    assert Path(repo.group(gid)['data']['project']).name.endswith('_v2')


def test_split_ready_part_exports_despite_other_part_translation_failure(setup,tmp_path):
    from transcomic.providers import KiraProvider
    repo,id,s,c=setup;cfg=repo.workset(id)['config'];cfg.update(capcut_mode='split',part_count=2,project_name='Story');repo.configure(id,cfg)
    ids=[add_group(repo,id,tmp_path/f'g{i}',f'{i+1}-{i+2}',1)[0] for i in range(4)]
    bad={r['id'] for gid in ids[:2] for r in repo.store.lines(repo.group(gid)['document_id'])}
    def handler(req):return response(req,[]) if any(r['id'] in bad for r in request_lines(req)) else response(req)
    s.translation.provider_factory=lambda store,job:KiraProvider('fixture','mock',transport=httpx.MockTransport(handler))
    run(c,id)
    projects=[a for a in repo.artifacts(id) if a['kind']=='capcut']
    assert len(projects)==1 and Path(projects[0]['path']).name=='Story_part2'
    assert [repo.group(gid)['data']['state'] for gid in ids]==['needs_attention','needs_attention','completed','completed']


def test_crash_metadata_recovery_does_not_run_any_engine(setup,tmp_path):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a');repo.activate(id)
    repo.execute('UPDATE automation_worksets SET owner_pid=?,owner_started=? WHERE id=?',(99999999,0,id))
    reopened=Repository(repo.store)
    assert reopened.workset(id)['state']=='interrupted' and reopened.groups(id)[0]['id']==gid
    assert FakeEngine.calls==[] and repo.store.jobs()==[]


def test_reopened_checkpoint_reuses_translations_and_verified_wavs(setup,tmp_path):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=6);run(c,id)
    meta=repo.group(gid)['data'];items=s.tts_store.items(meta['tts_run'])
    folder=Path(s.tts_store.docs(meta['tts_run'])[0]['folder'])
    good={p.name:file_hash(p) for p in folder.glob('*.wav') if p.name!='0004.wav'}
    (folder/'0004.wav').unlink();jobs=len(repo.store.jobs());before=len(FakeEngine.calls)
    repo.activate(id);repo.execute('UPDATE automation_worksets SET owner_pid=?,owner_started=? WHERE id=?',(99999999,0,id))
    # Independent connections reconstruct the checkpoint without live Python objects.
    fresh_store=Store(repo.store.path);fresh_tts=TtsStore(s.tts_store.path)
    try:
        fresh=Repository(fresh_store);fresh.tts_store=fresh_tts;fresh_tts.automation=fresh
        def forbidden(*args):raise AssertionError('Cached translations must not call HTTP')
        services=SimpleNamespace(tts_store=fresh_tts,translation=Translation(fresh,forbidden,network_probe=lambda provider:True),kokoro=Kokoro(fresh,fresh_tts,None,engine_factory=FakeEngine),capcut=Capcut(fresh,fresh_tts),preflight=lambda cfg:None)
        reopened=Controller(fresh,services)
        assert fresh.workset(id)['state']=='interrupted' and not reopened.busy
        assert len(FakeEngine.calls)==before and len(fresh_store.jobs())==jobs
        run(reopened,id)
        assert fresh.workset(id)['state']=='completed' and len(FakeEngine.calls)==before+1
        assert len(fresh_store.jobs())==jobs
        assert all(file_hash(folder/name)==digest for name,digest in good.items())
    finally:fresh_tts.close();fresh_store.close()


def test_tts_final_rescue_waits_for_other_ready_groups(setup,tmp_path,monkeypatch):
    repo,id,s,c=setup
    for n in range(2):add_group(repo,id,tmp_path/str(n),name=f'Chapter {n}',count=1)
    original=s.tts_store.claim;claims=[];first=[]
    def claim(*args,**kwargs):
        item=original(*args,**kwargs)
        if item:
            run_id=args[0]
            if not first:first.append(run_id)
            claims.append(run_id);FakeEngine.fail=lambda text:run_id==first[0]
        return item
    monkeypatch.setattr(s.tts_store,'claim',claim)
    run(c,id,timeout=20)
    assert claims.count(first[0])==5 and len(claims)==6
    assert claims[4]!=first[0] and claims[5]==first[0]
    assert len([a for a in repo.artifacts(id) if a['kind']=='capcut' and a['state']=='complete'])==1


def test_resume_selected_failed_group_does_not_run_other_group_or_next_workset(setup,tmp_path):
    repo,id,s,c=setup;a,_,_=add_group(repo,id,tmp_path/'a',count=1);b,_,_=add_group(repo,id,tmp_path/'b',count=1)
    next_id=repo.create('Next explicit workset',repo.workset(id)['config']);add_group(repo,next_id,tmp_path/'next',count=1)
    FakeEngine.fail=lambda text:True;run(c,id);before=len(FakeEngine.calls)
    FakeEngine.fail=None;assert c.validate(id)['passed'];c.start(id,selected=[a]);c.join(10)
    assert not c.busy and len(FakeEngine.calls)==before+1
    assert repo.group(a)['data']['state']=='completed' and repo.group(b)['data']['state']=='needs_attention'
    assert repo.workset(next_id)['state']=='draft' and not repo.workset(next_id)['started']


def test_failed_terminal_checkpoint_keeps_lock_until_explicit_recovery(setup,tmp_path,monkeypatch):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=1)
    original=repo.state
    def disk_full(*args):raise OSError(28,'Synthetic checkpoint disk full')
    monkeypatch.setattr(repo,'state',disk_full);run(c,id)
    assert not c.busy and c.pending_terminal=='needs_attention' and repo.active()==id
    with pytest.raises(ValueError):repo.configure(id,repo.workset(id)['config'])
    jobs=len(repo.store.jobs());calls=len(FakeEngine.calls)
    monkeypatch.setattr(repo,'state',original);assert c.validate(id)['passed']
    assert repo.active() is None and c.pending_terminal is None
    assert len(repo.store.jobs())==jobs and len(FakeEngine.calls)==calls


def test_manual_translation_rechecks_owner_inside_job_transaction(setup,tmp_path,monkeypatch):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a');original=repo.store.setting
    def activate_between_checks(key,default=''):
        if key=='retry_504' and not repo.active():repo.activate(id)
        return original(key,default)
    monkeypatch.setattr(repo.store,'setting',activate_between_checks)
    with pytest.raises(ValueError):repo.store.create_job(repo.group(gid)['document_id'],'mock','vi','en',1)
    assert repo.store.jobs()==[] and FakeEngine.calls==[]
    repo.state(id,'stopped')


def test_source_same_layout_retranslates_only_affected_batch_and_reports_are_private(setup,tmp_path):
    repo,id,s,c=setup;gid,path,_=add_group(repo,id,tmp_path/'a',count=9);run(c,id)
    before=sum(len(repo.store.attempts(b['id'])) for j in repo.store.jobs() for b in repo.store.batches(j['id']))
    lines=path.read_text(encoding='utf8').splitlines();lines[4]='Văn bản tổng hợp đã thay đổi';path.write_text('\n'.join(lines),encoding='utf8');FakeEngine.calls.clear();run(c,id)
    after=sum(len(repo.store.attempts(b['id'])) for j in repo.store.jobs() for b in repo.store.batches(j['id']))
    assert after-before==1
    # Mock translation text did not change; no voice regeneration is needed.
    assert FakeEngine.calls==[]
    reports=repo.workset(id)['checkpoint']['reports'];data=json.loads(Path(reports['json']).read_text(encoding='utf8'))
    assert data['state']=='completed' and data['groups'][0]['batches']
    assert 'automation-synthetic-key' not in json.dumps(data)
    assert '<html' in Path(reports['html']).read_text(encoding='utf8')


def test_add_or_remove_images_locks_only_affected_group(setup,tmp_path):
    repo,id,s,c=setup;a,_,folder=add_group(repo,id,tmp_path/'a','1-2');b,_,_=add_group(repo,id,tmp_path/'b','3-4');run(c,id)
    (folder/'0001.png').unlink();proof=c.validate(id);assert not proof['passed'] and not proof['errors']
    c.start(id);c.join(8)
    assert repo.group(a)['data']['state']=='needs_attention' and repo.group(b)['data']['state']=='completed'


def test_docx_and_numeric_slots_keep_document_structure_and_ids(setup,tmp_path):
    from docx import Document
    from transcomic.docx_document import import_docx
    repo,id,s,c=setup;source=tmp_path/'document.docx';doc=Document();doc.add_paragraph('Văn bản đầu');doc.add_paragraph('123');doc.add_paragraph('');doc.add_paragraph('...');doc.add_table(rows=1,cols=1).cell(0,0).text='Văn bản cuối';doc.save(source)
    original=file_hash(source);gid=repo.add(id,[source])[0];images=tmp_path/'images';images.mkdir()
    for n in range(3):
        image=QImage(20,20,QImage.Format_RGB32);image.fill(QColor('blue'));image.save(str(images/f'{n+1:04}.png'))
    repo.pair_images(gid,images);ids=[r['id'] for r in repo.store.lines(repo.group(gid)['document_id'])];run(c,id)
    assert repo.workset(id)['state']=='completed'
    assert len(FakeEngine.calls)==3 and '123' in FakeEngine.calls
    assert file_hash(source)==original and [r['id'] for r in repo.store.lines(repo.group(gid)['document_id'])]==ids
    result=import_docx(tmp_path/'document.en.docx');assert len(result.units)==len(import_docx(source).units)


def test_stop_during_inflight_http_does_not_dispatch_next_batch(setup,tmp_path):
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=6);cfg=repo.workset(id)['config'];cfg['workers']=1;repo.configure(id,cfg)
    entered=threading.Event();release=threading.Event();calls=[]
    async def handler(req):
        calls.append(request_lines(req));entered.set()
        while not release.is_set():await asyncio.sleep(.01)
        return response(req)
    s.translation.provider_factory=lambda store,job:KiraPoolProvider(store,'mock',transport=httpx.MockTransport(handler),rpm=100000,tpm=100000000)
    assert c.validate(id)['passed'];c.start(id);assert entered.wait(4);c.stop();assert repo.workset(id)['state']=='stopping'
    release.set();c.join(6)
    assert not c.busy and len(calls)==1 and FakeEngine.calls==[]
    assert sum(r['state']=='done' for r in repo.store.lines(repo.group(gid)['document_id']))==3
    run(c,id);assert repo.workset(id)['state']=='completed' and len(calls)==2


def test_quota_exhausted_stops_translation_but_ready_group_finishes(setup,tmp_path):
    repo,id,s,c=setup;a,_,_=add_group(repo,id,tmp_path/'a','1-2',1);b,_,_=add_group(repo,id,tmp_path/'b','3-4',1)
    cfg=repo.workset(id)['config'];cfg['workers']=1;repo.configure(id,cfg);calls=[]
    def handler(req):
        calls.append(request_lines(req))
        return response(req) if len(calls)==1 else httpx.Response(429,json={'error':{'code':'insufficient_quota','type':'insufficient_quota'}})
    s.translation.provider_factory=lambda store,job:KiraPoolProvider(store,'mock',transport=httpx.MockTransport(handler),rpm=100000,tpm=100000000)
    run(c,id)
    assert repo.group(a)['data']['state']=='completed' and repo.group(b)['data']['state']=='needs_attention'
    assert len(calls)==2 and len(FakeEngine.calls)==1


def test_network_probes_do_not_consume_request_attempts(setup,tmp_path):
    from transcomic.providers import KiraProvider
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=1);calls=[];probes=[]
    def handler(req):
        calls.append(1)
        if len(calls)==1:raise httpx.ConnectError('Synthetic offline',request=req)
        return response(req)
    def probe(provider):probes.append(provider);return len(probes)>=2
    s.translation.provider_factory=lambda store,job:KiraProvider('fixture','mock',transport=httpx.MockTransport(handler));s.translation.network_probe=probe
    run(c,id,timeout=12)
    assert len(calls)==2 and len(probes)==2 and repo.workset(id)['state']=='completed'
    batches=repo.group(gid)['data']['batch_jobs'];assert sum(len(repo.store.attempts(bid)) for value in batches.values() for bid in repo.batch_ids(value))==2


def test_old_database_tables_settings_and_manifest_migration_are_preserved(tmp_path):
    store=Store(tmp_path/'old.sqlite3')
    try:
        store.set_setting('model','existing-model');store.set_setting('tts_config','{"voice":"af_jessica","speed":1.1}');store.set_setting('capcut_workspace_v1','{"groups":[]}')
        tables=[r[0] for r in store.db.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        before={name:[tuple(r) for r in store.db.execute('SELECT * FROM "'+name+'"')] for name in tables}
        repo=Repository(store)
        assert {name:[tuple(r) for r in store.db.execute('SELECT * FROM "'+name+'"')] for name in tables}==before
        id=repo.create('Migration fixture',defaults(store));gid,path,_=add_group(repo,id,tmp_path/'g')
        meta=repo.group(gid)['data'];repo.execute('DELETE FROM automation_group_fields WHERE group_id=?',(gid,));repo.execute('UPDATE automation_groups SET data=? WHERE id=?',(json.dumps(meta),gid))
        upgraded=Repository(store);assert all(upgraded.group(gid)['data'][k]==v for k,v in meta.items())
        assert upgraded.group(gid)['data']['slot_count']==len(meta['slots'])
        assert len(upgraded.group_state(gid).get('images',[]))==0  # Large manifests migrated to field records.
    finally:store.close()


def test_reports_failure_keeps_successful_artifacts_and_internal_checkpoint(setup,tmp_path,monkeypatch):
    from transcomic.automation import reports
    repo,id,s,c=setup;gid,_,_=add_group(repo,id,tmp_path/'a',count=1)
    monkeypatch.setattr(reports,'write',lambda *a,**k:(_ for _ in ()).throw(PermissionError('Synthetic report denied')))
    run(c,id)
    assert repo.workset(id)['state']=='completed' and repo.group(gid)['data']['project']
    assert any(e['stage']=='report' for e in repo.events(id)) and repo.artifacts(id)


def test_project_error_is_local_other_independent_project_survives(setup,tmp_path,monkeypatch):
    from transcomic.capcut import exporter
    repo,id,s,c=setup;a,_,_=add_group(repo,id,tmp_path/'a','1-2',1);b,_,_=add_group(repo,id,tmp_path/'b','3-4',1);original=exporter.write_project
    def export(scans,output,name,*args,**kwargs):
        if name=='1-2':raise PermissionError('Synthetic permission denied')
        return original(scans,output,name,*args,**kwargs)
    monkeypatch.setattr(exporter,'write_project',export);run(c,id)
    assert repo.group(a)['data']['state']=='needs_attention' and repo.group(b)['data']['state']=='completed'
    assert not list((tmp_path/'capcut').glob('.transcomic-capcut-*'))
