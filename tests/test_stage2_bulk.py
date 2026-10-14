import asyncio
import json
import time
from pathlib import Path
import pytest
from docx import Document
from transcomic.bulk import import_many, discover_files, plan_exports, write_exports, workspace_status
from transcomic.documents import decode_txt
from transcomic.providers import BatchResult
from transcomic.storage import Store
from transcomic.translation import Scheduler
import httpx
from transcomic.gemini import GeminiProvider


@pytest.mark.parametrize('count', [1, 4])
def test_model_404_finishes_empty_queue_or_pauses_unsent_work(tmp_path, count):
    store = Store(tmp_path/'state.db')
    source = tmp_path/'source.txt'; source.write_text('\n'.join('Line '+str(i) for i in range(count)))
    doc, _ = store.import_document(source, decode_txt(source.read_bytes()), 'vi', 'en')
    job = store.create_job(doc, 'mock', 'vi', 'en', 1, provider='gemini')
    calls = []
    def handle(req): calls.append(req); return httpx.Response(404, json={'error': {'status': 'NOT_FOUND'}})
    provider = GeminiProvider([{'id': 'key', 'secret': 'secret', 'project': 'p'}], 'mock', transport=httpx.MockTransport(handle))
    scheduler = Scheduler(store, provider); scheduler.start(job)
    scheduler.join(2)
    assert not scheduler._paused.is_set() and len(calls) == 1
    assert not scheduler._thread.is_alive()
    assert scheduler.intervention_reason == 'model_not_found'
    assert store.job(job)['state'] == 'needs_intervention'
    if count == 4: assert len(store.batches(job, 'pending')) == 1
    store.close()


@pytest.mark.parametrize('count', [1, 10, 100])
def test_nested_import_natural_order_duplicates_bad_file_and_reimport(tmp_path, count):
    source = tmp_path/'source'; source.mkdir(); (source/'nested').mkdir()
    for i in range(count): (source/'nested'/f'Chapter {i+1}.txt').write_text('First\nSecond', encoding='utf-8')
    word = Document(); word.add_paragraph('Word source'); word.save(source/'Word.docx')
    (source/'bad.docx').write_bytes(b'not a zip')
    (source/'unsupported.srt').write_text('unsupported')
    store = Store(tmp_path/'state.db')
    first = import_many(store, [source, source/'nested'/'Chapter 1.txt'], 'Tiếng Việt', 'Tiếng Anh')
    assert len(first['document_ids']) == count+1 and len(first['errors']) == 2
    names = [Path(store.document(d)['path']).name for d in first['document_ids']]
    if count >= 10: assert names.index('Chapter 2.txt') < names.index('Chapter 10.txt')
    second = import_many(store, [source/'nested'/'Chapter 1.txt'], 'Tiếng Việt', 'Tiếng Anh')
    assert second['document_ids'][0] not in first['document_ids']
    assert all(r['translation'] is None for d in first['document_ids'] for r in store.lines(d))
    assert not store.jobs()
    store.close()


class Fake:
    def __init__(self): self.calls = []; self.active = self.peak = 0
    async def translate(self, lines, source, target):
        self.active += 1; self.peak = max(self.peak, self.active)
        self.calls.append((target, [r['id'] for r in lines], [c['source'] for r in lines for c in r.get('context', [])]))
        await asyncio.sleep(.02)
        self.active -= 1
        return BatchResult({r['id']: target+' '+r['source'] for r in reversed(lines)})


@pytest.mark.parametrize('concurrency', [1, 3, 10, 30, 40, 50])
def test_global_workers_fair_languages_files_no_cross_context(tmp_path, concurrency):
    store = Store(tmp_path/'state.db')
    ids = [store.import_document(tmp_path/f'{prefix}.txt', decode_txt('\n'.join(prefix+str(i) for i in range(90)).encode()), 'Tiếng Việt', 'Tiếng Anh')[0] for prefix in ('AAA', 'BBB')]
    run, jobs = store.create_run(ids, ['Tiếng Anh', 'Tiếng Indonesia'], provider='kira', model='frozen', source_lang='Tiếng Việt', concurrency=concurrency)
    fake = Fake(); scheduler = Scheduler(store, fake)
    scheduler.start_many(jobs, run); scheduler.join(15)
    assert not scheduler._thread.is_alive() and fake.peak == concurrency
    first_jobs = [next(j for j in jobs if call[1][0] in {r['id'] for r in store.lines(store.job(j)['document_id'])}) for call in fake.calls[:4]]
    assert len(set(first_jobs)) == 2  # Fair between files in the current language.
    assert [c[0] for c in fake.calls]==['Tiếng Anh']*60+['Tiếng Indonesia']*60
    assert len(fake.calls) == 120
    for job in store.run_jobs(run):
        rows = store.lines(job['document_id'])
        assert all(r['translation'].startswith(job['target_lang']) for r in rows)
        expected_prefix = rows[0]['source'][:3]
        for target, lineids, context in fake.calls:
            if lineids[0] in {r['id'] for r in rows}: assert all(c.startswith(expected_prefix) for c in context)
    assert store.run(run)['state'] == 'finished'
    store.close()


@pytest.mark.parametrize('source,targets', [('Tiếng Việt', ['Tiếng Anh']), ('Tiếng Việt', ['Tiếng Indonesia']), ('Tiếng Việt', ['Tiếng Anh', 'Tiếng Indonesia']), ('Tiếng Anh', ['Tiếng Việt']), ('Tiếng Indonesia', ['Tiếng Việt'])])
def test_pair_workspaces_independent_ids_edits_errors_history(tmp_path, source, targets):
    store = Store(tmp_path/'state.db')
    doc, _ = store.import_document(tmp_path/'a.txt', decode_txt(b'Source'), source, targets[0])
    run, jobs = store.create_run([doc], targets, provider='kira', model='model', source_lang=source, concurrency=1)
    assert len(jobs) == len(targets)
    ids = []
    for job in store.run_jobs(run):
        row = store.lines(job['document_id'])[0]; ids.append(row['id'])
        store.edit_line(row['id'], 'translation', job['target_lang'])
    assert len(set(ids)) == len(ids)
    assert len(store.history()) == len(targets)
    assert store.document(doc)['snapshot'] == b'Source'
    store.close()


def test_source_same_target_and_empty_run_roll_back(tmp_path):
    store = Store(tmp_path/'state.db')
    doc, _ = store.import_document(tmp_path/'a.txt', decode_txt(b'Data'))
    with pytest.raises(ValueError): store.create_run([doc], ['Tiếng Việt'], provider='kira', model='m', source_lang='Tiếng Việt', concurrency=3)
    assert not store.jobs() and store.db.execute('SELECT COUNT(*) FROM translation_runs').fetchone()[0] == 0
    store.close()


def test_export_stable_collision_snapshot_and_versions(tmp_path):
    store = Store(tmp_path/'state.db'); ids = []
    for folder in ('first', 'second'):
        path = tmp_path/folder/'Chapter.txt'; path.parent.mkdir(); path.write_bytes(b'Source\r\n\r\nSecond')
        doc, _ = store.import_document(path, decode_txt(path.read_bytes()), 'Tiếng Việt', 'Tiếng Anh'); ids.append(doc)
    run, jobs = store.create_run(ids, ['Tiếng Anh', 'Tiếng Indonesia'], provider='kira', model='m', source_lang='Tiếng Việt', concurrency=1)
    docs = [store.job(j)['document_id'] for j in jobs]
    for doc in docs: store.edit_line(store.lines(doc)[0]['id'], 'translation', 'Captured')
    captured = store.export_snapshot(docs)
    plan = plan_exports(captured, tmp_path/'out')
    reverse = plan_exports(list(reversed(captured)), tmp_path/'out')
    assert [p['target'] for p in plan] == [p['target'] for p in reverse]
    assert len({str(p['target']) for p in plan}) == 4
    store.edit_line(store.lines(docs[0])[0]['id'], 'translation', 'Late result')
    results = write_exports(store, plan, allow_incomplete=True)
    assert all(r['state'] == 'exported' for r in results)
    assert all(Path(r['path']).read_bytes() == b'Captured\r\n\r\nSecond' for r in results)
    assert 'Late result' not in Path(results[0]['path']).read_text()
    assert all(r['state'] == 'cancelled' for r in write_exports(store, plan, existing='cancel', allow_incomplete=True))
    versions = write_exports(store, plan, existing='version', allow_incomplete=True)
    assert all(Path(r['path']).stem.endswith('_2') for r in versions)
    assert all((tmp_path/folder/'Chapter.txt').read_bytes() == b'Source\r\n\r\nSecond' for folder in ('first', 'second'))
    store.close()


def test_docx_bulk_export_structure_and_incomplete_source(tmp_path):
    path = tmp_path/'a.docx'; docx = Document(); docx.add_paragraph('Paragraph'); table = docx.add_table(rows=1, cols=1); table.cell(0, 0).text = 'Cell'; docx.save(path)
    original = path.read_bytes(); store = Store(tmp_path/'state.db')
    ids = import_many(store, [path], 'Tiếng Việt', 'Tiếng Anh')['document_ids']
    run, jobs = store.create_run(ids, ['Tiếng Anh', 'Tiếng Indonesia'], provider='kira', model='m', source_lang='Tiếng Việt', concurrency=1)
    workspaces = [store.job(j)['document_id'] for j in jobs]
    for d in workspaces: store.edit_line(store.lines(d)[0]['id'], 'translation', store.document(d)['target_lang'])
    plan = plan_exports(store.export_snapshot(workspaces))
    assert all('TransComic' in p['target'].parts for p in plan)
    assert all(r['state'] == 'error' for r in write_exports(store, plan))
    results = write_exports(store, plan, allow_incomplete=True)
    for result in results:
        output = Document(result['path']); assert len(output.tables) == 1 and output.tables[0].cell(0,0).text == 'Cell'
        assert output.paragraphs[0].text in ('Tiếng Anh', 'Tiếng Indonesia')
    assert path.read_bytes() == original
    store.close()


def test_run_crash_recovery_pending_only_no_replay(tmp_path):
    store = Store(tmp_path/'state.db'); doc, _ = store.import_document(tmp_path/'a.txt', decode_txt(b'A\nB\nC\nD'))
    run, jobs = store.create_run([doc], ['Tiếng Anh', 'Tiếng Indonesia'], provider='kira', model='m', source_lang='Tiếng Việt', concurrency=1)
    store.run_state(run, 'running'); store.job_state(jobs[0], 'running')
    batch = store.batches(jobs[0])[0]; store.claim_batch(batch['id']); store.prepare_attempt(batch['id'], 1)
    store.close(); store = Store(tmp_path/'state.db')
    assert store.run(run)['state'] == 'interrupted' and store.batches(jobs[0])[0]['state'] == 'unknown'
    assert len(store.attempts(batch['id'])) == 1
    fake = Fake(); scheduler = Scheduler(store, fake); scheduler.start_many(jobs, run); scheduler.join(5)
    assert batch['payload'][0]['id'] not in {i for c in fake.calls for i in c[1]}
    assert any(r['state'] == 'unknown' for r in store.lines(doc))
    assert all(c[0]=='Tiếng Anh' for c in fake.calls)
    assert store.run(run)['state']=='needs_intervention'
    assert store.run(run)['config']['current_locale']=='en'
    store.close()


def test_global_request_budget_counts_all_jobs(tmp_path):
    store = Store(tmp_path/'state.db'); doc, _ = store.import_document(tmp_path/'a.txt', decode_txt(b'A\nB\nC\nD'))
    run, jobs = store.create_run([doc], ['Tiếng Anh', 'Tiếng Indonesia'], provider='kira', model='m', source_lang='Tiếng Việt', concurrency=3, config={'max_requests': 1})
    fake = Fake(); scheduler = Scheduler(store, fake); scheduler.start_many(jobs, run); scheduler.join(5)
    assert len(fake.calls) == 1
    assert store.db.execute('SELECT COUNT(*) FROM request_attempts').fetchone()[0] == 1
    store.close()
