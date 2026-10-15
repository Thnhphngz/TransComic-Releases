from pathlib import Path
import json

import httpx
import pytest
from docx import Document

from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox, QInputDialog
from PySide6.QtTest import QTest

from transcomic.documents import decode_txt
from transcomic.providers import KiraProvider
from transcomic.storage import Store
from transcomic.translation import Scheduler
from transcomic.ui.main import MainWindow


def test_import_edit_export_through_window(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "chapter.txt"
    source.write_bytes(" Hello \r\n\r\nWorld".encode())
    output = tmp_path / "chapter_VI.txt"
    store = Store(tmp_path / "app.db")
    window = MainWindow(store)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(source), ""))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(output), ""))
    monkeypatch.setattr(QMessageBox, "exec", lambda self: self.buttons()[0].click())
    window.import_file()
    assert window.model.rowCount() == 3
    assert window.model.setData(window.model.index(0, 1), "Xin chào")
    assert window.model.setData(window.model.index(2, 1), "Thế giới")
    monkeypatch.setattr(QInputDialog, 'getItem', lambda *args: (args[3][0], True))
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.Yes)
    window.export_file()
    for _ in range(200):
        app.processEvents()
        if not window.export_busy: break
        QTest.qWait(10)
    output = source.with_name('chapter_EN.txt')
    assert output.read_bytes() == " Xin chào \r\n\r\nThế giới".encode()
    assert source.read_bytes() == b" Hello \r\n\r\nWorld"
    window.close()


@pytest.mark.parametrize("count", [10, 100, 1000, 10000])
def test_editor_expands_scrolls_and_keeps_position(tmp_path: Path, monkeypatch, count: int):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    store = Store(tmp_path / "large.db")
    text = "\n".join("Dòng tiếng Việt %d" % i if i % 7 else "" for i in range(count))
    doc_id, _ = store.import_document(tmp_path / "chapter.txt", decode_txt(text.encode()))
    window = MainWindow(store)
    window.resize(1600, 1000)
    window.show()
    window.open_document(doc_id)
    app.processEvents()
    table = window.table
    assert table.viewport().height() >= 500
    assert table.rowAt(0) == 0
    assert table.visualRect(window.model.index(min(count - 1, 14), 0)).bottom() < table.viewport().height()
    table.scrollToBottom()
    app.processEvents()
    assert table.visualRect(window.model.index(count - 1, 0)).bottom() < table.viewport().height()
    top_before = table.rowAt(0)
    visible = store.lines(doc_id)[top_before]
    assert window.model.setData(window.model.index(top_before, 1), "Đã sửa")
    assert store.line(visible["id"])["translation"] == "Đã sửa"
    assert table.rowAt(0) == top_before
    changed = store.lines(doc_id)[max(0, count // 2)]
    store.edit_line(changed["id"], "translation", "API mock")
    window.refresh([changed["id"]])
    app.processEvents()
    assert table.rowAt(0) == top_before
    window.close()


def test_docx_gui_import_resume_mock_translate_export(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "chapter.docx"
    original = Document()
    p = original.add_paragraph()
    p.add_run("Hello ")
    p.add_run("world").bold = True
    p = original.add_paragraph()
    p.add_run("First")
    p.add_run().add_break()
    p.add_run("Second")
    table = original.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    original.save(source)
    store = Store(tmp_path / "docx_gui.db")
    window = MainWindow(store)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(source), ""))
    window.import_file()
    doc_id = window.doc_id
    assert doc_id and store.document(doc_id)["format"] == "docx"
    assert any(row["source"] == "Cell" for row in store.units(doc_id))
    window.close()
    source.unlink()
    reopened_store = Store(tmp_path / "docx_gui.db")
    reopened = MainWindow(reopened_store)
    assert reopened.doc_id is None
    assert reopened.model.rowCount() == 0
    reopened.open_document(doc_id)  # Old snapshots are opened explicitly after empty startup.

    def handler(request):
        body = json.loads(request.content)
        assert body["model"] == "qwen3.8-flash"
        payload = json.loads(body["messages"][1]["content"])
        lines = [{"id": item["id"], "text": "VI " + item["text"]} for item in payload["lines"]]
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"lines": lines})}}]})

    job_id = reopened_store.create_job(doc_id, "qwen3.8-flash", "Ti\u1ebfng Anh", "Ti\u1ebfng Vi\u1ec7t", 2)
    scheduler = Scheduler(reopened_store, KiraProvider("fake", "qwen3.8-flash", transport=httpx.MockTransport(handler)))
    scheduler.start(job_id)
    scheduler.join(10)
    app.processEvents()
    assert all(row["state"] in ("done", "skip") for row in reopened_store.units(doc_id))
    edited = next(row for row in reopened_store.units(doc_id) if row["source"] == "Cell")
    reopened_store.edit_line(edited["id"], "translation", "Ô đã sửa")
    output = tmp_path / "chapter_VI.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(output), ""))
    monkeypatch.setattr(QMessageBox, "exec", lambda self: self.buttons()[0].click())
    monkeypatch.setattr(QInputDialog, 'getItem', lambda *args: (args[3][2], True))
    monkeypatch.setattr(QFileDialog, 'getExistingDirectory', lambda *args: str(tmp_path))
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: QMessageBox.Yes)
    reopened.export_file()
    for _ in range(200):
        app.processEvents()
        if not reopened.export_busy: break
        QTest.qWait(10)
    output = tmp_path / 'chapter_VI.docx' 
    rendered = Document(output)
    assert rendered.paragraphs[0].text == "VI Hello world"
    assert rendered.paragraphs[1].text == "VI First\nVI Second"
    assert rendered.tables[0].cell(0, 0).text == "Ô đã sửa"
    assert source.exists() is False
    reopened.close()


def test_long_rows_wrap_without_creating_new_lines(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    store = Store(tmp_path / "long.db")
    source = "Đây là một câu rất dài. " * 100
    doc_id, _ = store.import_document(tmp_path / "long.txt", decode_txt((source + "\nNgắn").encode()))
    window = MainWindow(store)
    window.resize(1200, 720)
    window.show()
    window.open_document(doc_id)
    app.processEvents()
    assert window.model.rowCount() == 2
    assert window.table.rowHeight(0) > window.table.rowHeight(1)
    QTest.qWait(30)  # Allow Qt's deferred scroll geometry to settle after row resizing.
    assert window.table.verticalScrollBar().isVisible()
    window.resize(900, 580)
    app.processEvents()
    assert window.table.viewport().height() > 0
    def editor_rect(widget):
        from PySide6.QtCore import QRect
        return QRect(widget.mapTo(window.editor_page, widget.rect().topLeft()), widget.size())
    assert editor_rect(window.export_button).right() < editor_rect(window.export_many_button).left()
    assert editor_rect(window.export_many_button).right() < editor_rect(window.folder_button).left()
    assert editor_rect(window.folder_button).right() < window.editor_page.width()
    assert window.width() == 900
    window.close()
