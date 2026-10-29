"""Policy routing and DOCX/TTS contracts. Mock text is not AI quality evidence."""
import asyncio
import io
import json
from pathlib import Path
import httpx
import pytest
from docx import Document
from transcomic.prompts import messages,PROMPT_VERSION
from transcomic.translation_policy import VERSION
from transcomic.quality import warnings
from transcomic.documents import decode_txt,ImportedText,TextLine
from transcomic.storage import Store
from transcomic.deepseek import DeepSeekProvider,MODEL,save_key,provider_for
from transcomic.translation import Scheduler
from transcomic.docx_document import inspect_docx,render_docx
from transcomic.tts.documents import read_document,segments


def test_shared_semantic_rules_have_no_punctuation_forcing():
    lines=[dict(id='stable',source='Đây là câu hỏi sao',semantic_facts=[dict(fact='Soles is addressed as a boy in this passage',evidence=['source:58'])])]
    new=messages(lines,'vi','en',VERSION);old=messages(lines,'vi','en',PROMPT_VERSION)
    assert 'Punctuation and quotation marks may be added' in new[0]['content']
    assert 'Use no internal or final periods' not in new[0]['content']
    assert 'Use no internal or final periods' in old[0]['content']
    assert json.loads(new[1]['content'])['verified_story_facts'] and 'stable' in new[0]['content']


@pytest.mark.parametrize('version,expected',[(VERSION,'Is he here? Yes. "Welcome!"'),(PROMPT_VERSION,'Is he here? Yes, "Welcome!"')])
def test_new_policy_preserves_api_text_old_snapshot_keeps_legacy(tmp_path,version,expected):
    store=Store(tmp_path/'s.db');save_key(store,'synthetic')
    doc,_=store.import_document(tmp_path/'text.txt',decode_txt('Anh ấy ở đây sao'.encode()),'vi','en')
    job=store.create_job(doc,MODEL,'vi','en',1,provider='deepseek',config={'prompt_version':version})
    def handler(r):
        rows=json.loads(json.loads(r.content)['messages'][1]['content'])['lines']
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'lines':[{'id':rows[0]['id'],'text':'Is he here? Yes. "Welcome!"'}]})}}]})
    p=provider_for(store,store.job(job));p.transport=httpx.MockTransport(handler);s=Scheduler(store,p)
    try:
        s.start(job);s.join(3)
        assert store.lines(doc)[0]['translation']==expected
        assert store.job(job)['prompt_version']==version
        if version==VERSION:
            assert store.line_policies(doc)[store.lines(doc)[0]['id']]==VERSION
            assert not any(t in ('Thêm dấu chấm','Dấu câu cuối') for r in store.warning_rows(doc) for t,n in r['warnings'])
    finally:store.close()


def test_semantic_qc_uses_question_meaning_not_question_mark_only():
    row=dict(source='Những bức tranh này do chị Lan vẽ sao',translation='Were these pictures painted by Lan',state='done',source_lang='vi',target_lang='en',translation_policy=VERSION)
    assert not any(k in ('Sắc thái câu hỏi','Dấu câu cuối','Thêm dấu chấm') for k,v in warnings(row))
    row['translation']='These pictures were painted by Lan.'
    assert any(k=='Sắc thái câu hỏi' for k,v in warnings(row))
    row.update(source='Anh ấy không phải người giao hàng',translation='He is the delivery person.')
    assert any(k=='Phủ định' for k,v in warnings(row))


@pytest.mark.parametrize('source,target,kind',[
 ('sự hỗn loạn chớt chóc bao trùm lên tất cả','A sudden chaos descended upon everything','Thiếu ý nghĩa'),
 ('ai cũng gào lên, bắt lấy con sói con đó','Everyone was shouting, grabbing at that little wolf','Câu mệnh lệnh'),
])
def test_real_pilot_regressions_warn_without_rewriting(source,target,kind):
    row=dict(source=source,translation=target,state='done',source_lang='vi',target_lang='en',translation_policy=VERSION)
    assert any(k==kind for k,v in warnings(row))
    assert row['translation']==target


@pytest.mark.parametrize('separator',['\n','\r','\u0085','\u2028','\u2029'])
def test_semantic_line_breaks_rejected(separator):
    def handler(r):return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'lines':[dict(id='x',text='First'+separator+'second')]})}}]})
    p=DeepSeekProvider('fixture',transport=httpx.MockTransport(handler))
    result=asyncio.run(p.translate([dict(id='x',source='First')],'vi','en'))
    assert result.error and not result.values


def test_provider_defaults_and_kira_gemini_policy_unchanged(tmp_path):
    store=Store(tmp_path/'s.db')
    try:
        for provider,model,version in [('deepseek',MODEL,VERSION),('kira','mock',PROMPT_VERSION),('gemini','models/mock',PROMPT_VERSION)]:
            doc,_=store.import_document(tmp_path/f'{provider}.txt',decode_txt(b'source'),'vi','en')
            job=store.create_job(doc,model,'vi','en',1,provider=provider)
            assert store.job(job)['prompt_version']==version
    finally:store.close()


def test_docx_cache_roundtrip_long_unicode_blank_and_tts(tmp_path):
    d=Document();d.add_paragraph('Nguồn có chữ Việt '+('rất dài '*2000));d.add_paragraph('');d.add_paragraph('Câu hỏi sao')
    raw=io.BytesIO();d.save(raw);raw=raw.getvalue();a=inspect_docx(raw);b=inspect_docx(raw)
    assert a is b and len(a.units)==3 and a.units[1].text==''
    text=['A long English sentence. Another clause, kept in the same logical line.','','Is that a question?']
    rows=[dict(id=str(i),document_id='doc',ordinal=i,locator=u.locator,source=u.text,translation=text[i] or None,state='done' if text[i] else 'skip',source_revision=0,translation_source_revision=0) for i,u in enumerate(a.units)]
    output,missing=render_docx(raw,rows);assert missing==0
    from transcomic.documents import preserve_edge_space
    reread=inspect_docx(output);assert [u.text for u in reread.units]==[preserve_edge_space(u.text,t) for u,t in zip(a.units,text)]
    file=tmp_path/'out.docx';file.write_bytes(output)
    parsed=read_document(file);assert parsed['line_count']==3
    tts=segments(parsed['text'],'lines');assert len(tts)==2 and [r['line_start'] for r in tts]==[1,3]
    assert len(segments(text[0],'lines'))==1
    changed=Document(io.BytesIO(raw));changed.paragraphs[0].text='Changed';v=io.BytesIO();changed.save(v)
    assert inspect_docx(v.getvalue()).units[0].text=='Changed'
