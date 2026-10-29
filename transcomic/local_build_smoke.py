"""Opt-in frozen devtest smoke: isolated data and HTTP mocks only, never install."""
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import traceback
import httpx
from PySide6.QtCore import QTimer,Qt,QEventLoop,QPoint
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication
from .storage import Store
from .documents import decode_txt
from .ui.main import MainWindow
from .update.internal import InternalUpdateService,manifest
from .update.metadata import canonical,asset_url,release_url
from .update.service import LATEST_URL
from .version import VERSION


def mock_service(directory):
    from .release_version import bumped
    version=bumped(VERSION);tag='v'+version;name=f'TransComic-Setup-{version}-windows-x64.testdata';data=b'nonexecutable update fixture'
    envelope=manifest(version,name.replace('.testdata','.exe'),len(data),hashlib.sha256(data).hexdigest(),'2026-10-14T00:00:00Z')
    envelope['payload']['installer_filename']=name;raw=canonical(envelope)
    contents={name:data,'transcomic-update.json':raw,'SHA256SUMS.txt':(hashlib.sha256(data).hexdigest()+'  '+name+'\n'+hashlib.sha256(raw).hexdigest()+'  transcomic-update.json\n').encode()}
    release={'tag_name':tag,'draft':False,'prerelease':False,'html_url':release_url(tag),'published_at':'2026-10-14T00:00:00Z','body':'Internal update smoke <script>text only</script>',
             'assets':[{'name':n,'size':len(b),'state':'uploaded','browser_download_url':asset_url(tag,n)} for n,b in contents.items()]}
    def respond(req):
        return httpx.Response(200,json=release) if str(req.url)==LATEST_URL else httpx.Response(200,content=contents[req.url.path.rsplit('/',1)[-1]])
    return InternalUpdateService(directory,httpx.Client(transport=httpx.MockTransport(respond)),testing=True)


def main():
    root=Path(os.environ.get('LOCALAPPDATA','')).resolve()
    from .app import data_path
    from .version import APP_ID
    from .update.windows_helper import INSTALL_APP_ID
    if data_path()!=root/'TransComic/transcomic.sqlite3' or APP_ID!='TransComic.Desktop' or INSTALL_APP_ID!='{CB58C770-47B7-42DC-A2CA-409DDF456487}':return 2
    # This mode cannot open the normal user database, even when called without isolation.
    if not root.is_relative_to(Path(tempfile.gettempdir()).resolve()) or not root.name.startswith('TransComic-build-smoke-') or not (root/'isolated-smoke.json').is_file():return 2
    if json.loads((root/'isolated-smoke.json').read_text())!={'mode':'isolated-build-smoke'}:return 2
    original=httpx.Client.send
    def guard(client,*args,**kwargs):
        if not isinstance(client._transport,httpx.MockTransport):raise AssertionError('Real HTTP forbidden')
        return original(client,*args,**kwargs)
    httpx.Client.send=guard
    original_async=httpx.AsyncClient.send
    async def guard_async(client,*args,**kwargs):
        if not isinstance(client._transport,httpx.MockTransport):raise AssertionError('Real HTTP forbidden')
        return await original_async(client,*args,**kwargs)
    httpx.AsyncClient.send=guard_async
    app=QApplication(sys.argv);app.setApplicationName('TransComic');app.setAttribute(Qt.AA_DontUseNativeDialogs,True)
    if os.environ.get('TRANSCOMIC_SMOKE_CPU_ONLY')=='1':
        from .release_smoke_203 import cpu_probe
        try:proof=cpu_probe(root);code=0
        except Exception:proof=dict(passed=False,status='NOT_VERIFIED',error=traceback.format_exc());code=1
        (root/'cpu-probe.json').write_text(json.dumps(proof,indent=2),encoding='utf8')
        return code
    # This opt-in test never discovers/downloads a runtime or probes a GPU.
    from .ui.tts_page import TtsPage
    TtsPage.check_runtime=lambda self:None
    store=Store(root/'TransComic/transcomic.sqlite3');store.set_setting('update_auto_check','0')
    window=MainWindow(store);window.update_panel.auto.setChecked(False)
    window.show();app.setQuitOnLastWindowClosed(False)
    def wait(condition,seconds=5):
        if condition():return
        loop=QEventLoop();poll=QTimer();poll.setInterval(20);poll.setTimerType(Qt.PreciseTimer)
        poll.timeout.connect(lambda:loop.quit() if condition() else None)
        deadline=QTimer();deadline.setSingleShot(True);deadline.setTimerType(Qt.PreciseTimer);deadline.timeout.connect(loop.quit)
        poll.start();deadline.start(int(seconds*1000))
        try:loop.exec()
        finally:poll.stop();deadline.stop()
        assert condition(),'Smoke condition timeout'
    def exercise():
        try:
            assert store.path.resolve().is_relative_to(root)
            from .release_smoke import verify_baseline,run_automation,verify_reopen,verify_gateway_recovery
            if os.environ.get('TRANSCOMIC_SMOKE_REOPEN')=='1':
                verify_reopen(window,root)
                window.confirm_exit=lambda running:'safe';window.close()
                QTimer.singleShot(900,app.quit)
                return
            migration=verify_baseline(window,root)
            recovery=verify_gateway_recovery(window,root,wait)
            window.workspace_ids=[]
            from docx import Document
            from .docx_document import render_docx,inspect_docx
            text=root/'fixture.txt';text.write_bytes(b'Alpha one\nAlpha two')
            word=root/'fixture.docx';generated=Document();generated.add_paragraph('Alpha one');generated.add_paragraph('Alpha two');generated.save(word)
            original_hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (text,word)}
            window.import_paths([str(text),str(word)]);wait(lambda:not window.import_busy)
            assert len(window.workspace_ids)==2
            doc=next(i for i in window.workspace_ids if store.document(i)['format']=='txt')
            word_id=next(i for i in window.workspace_ids if store.document(i)['format']=='docx')
            for row in store.units(word_id):store.edit_line(row['id'],'translation','Translated '+row['source'])
            rendered,missing=render_docx(word.read_bytes(),store.units(word_id))
            assert missing==0 and len(inspect_docx(rendered).units)==2
            window.open_document(word_id);QTest.qWait(40);window.grab().save(str(root/'docx-fixture.png'))
            window.open_document(doc);window.model.setData(window.model.index(0,1),'Alpha translation')
            QTest.qWait(300);QApplication.processEvents()
            before=store.lines(doc);targets=window.target_languages()
            for dark in (False,True):
                if window.dark!=dark:window.toggle_theme()
                for compact in (False,True):
                    if (window.sidebar.width()<100)!=compact:window.toggle_sidebar()
                    if compact:window.showMaximized()
                    else:window.showNormal();window.resize(900,580)
                    window.pages.setCurrentIndex(9);QTest.qWait(150)
                    assert window.navigation_group.checkedId()==9
                    assert [window.navigation_group.id(b) for b in window.nav_buttons]==[0,7,8,9,1,2,3]
                    # Navigation now lives inside its own scroll viewport;
                    # compare coordinates in the sidebar, not different parents.
                    last=window.nav_buttons[-1];window.navigation_scroll.ensureWidgetVisible(last)
                    QTest.qWait(40)
                    assert window.navigation_scroll.viewport().rect().contains(last.mapTo(window.navigation_scroll.viewport(),last.rect().center()))
                    assert last.mapTo(window.sidebar,QPoint(0,last.height())).y()<=window.collapse_button.y()
                    QTest.mouseClick(last,Qt.LeftButton);assert window.pages.currentIndex()==3
                    window.pages.setCurrentIndex(9)
                    assert window.sky_switch.geometry().bottom()<=window.sidebar.height()
                    window.grab().save(str(root/f'automation-{dark}-{compact}.png'))
                    window.pages.setCurrentIndex(0);assert store.lines(doc)==before and window.target_languages()==targets,repr({'lines_before':before,'lines_after':store.lines(doc),'targets_before':targets,'targets_after':window.target_languages()})
                window.search.open(True);window.search.scope.setCurrentIndex(2);window.search.query.setText('Alpha');window.search.replacement.setText('Beta');window.search.navigate(1);window.search.replace_current()
                assert store.lines(doc)[0]['translation']=='Beta translation';window.search.undo();window.search.hide()
                before=store.lines(doc)  # Replace/Undo deliberately advances edit revisions.
            job=store.create_job(doc,'fixture','Tiếng Việt','Tiếng Anh',1);store.job_state(job,'completed')
            window.pages.setCurrentIndex(1);window.refresh();window.confirm_clear_all=lambda count:True;window.clear_all_queue()
            wait(lambda:not window.queue_clear_busy);assert not store.history() and store.lines(doc)==before
            panel=window.update_panel;panel.service=mock_service(root/'Updates');window.pages.setCurrentIndex(3)
            panel.check_now(False);wait(lambda:not panel.busy);assert panel.update_dialog is None
            for dark in (False,True):
                if window.dark!=dark:window.toggle_theme()
                panel.check_now(True);wait(lambda:panel.update_dialog is not None)
                panel.update_dialog.grab().save(str(root/f'popup-{dark}.png'));panel.update_dialog.reject();wait(lambda:panel.update_dialog is None)
                assert panel.downloaded_path is None
            panel.check_now(True);wait(lambda:panel.update_dialog is not None);panel.update_dialog.download.click()
            wait(lambda:not panel.busy and panel.downloaded_path is not None)
            assert panel.downloaded_path.suffix=='.testdata'
            window.pages.setCurrentIndex(7);QApplication.processEvents();assert hasattr(window,'tts_page')
            from .tts.managed_runtime import catalog,WORKER_FILES
            assert catalog()['python'] and all((Path(__file__).parent/'tts'/('worker_source/'+n if getattr(sys,'frozen',False) else n)).is_file() for n in WORKER_FILES)
            tts={'page_open':True,'catalog_and_worker_paths':True,'cached_cpu_load':'NOT_VERIFIED'}
            cpu_cache=os.environ.get('TRANSCOMIC_SMOKE_CPU_CACHE')
            if cpu_cache:
                cache=Path(cpu_cache).resolve();marker=json.loads((cache/'isolated-distribution.json').read_text(encoding='utf8'))
                assert cache.name.startswith('TransComic-distribution-cpu-') and marker.get('backend')=='cpu'
                from .tts.models import ModelManager
                from .tts.managed_runtime import ManagedRuntimes
                from .tts.managed_engine import ManagedEngine
                from .tts.audio import encode
                models=ModelManager(cache/'TransComic/TTS/Models');assert models.ready('standard')
                runtime=ManagedRuntimes(cache/'TransComic/TTS/Runtimes').active();assert runtime['backend']=='CPUExecutionProvider'
                engine=ManagedEngine(models,dict(variant='standard',language='en-us',voice='af_heart',speed=1),runtime,width=1)
                try:
                    engine.probe();blocks=list(engine.synthesize('Hello.', 'af_heart',1,'en-us'))
                    wav=encode(root/'kokoro-probe.wav',iter(blocks),'wav');mp3=encode(root/'kokoro-probe.mp3',iter(blocks),'mp3')
                    tts.update(cached_cpu_load='PASS',audio_wav=bool(wav),audio_mp3=bool(mp3),backend=runtime['backend'],download_bytes=0)
                finally:engine.close()
            window.pages.setCurrentIndex(8);QApplication.processEvents();assert hasattr(window,'capcut_page')
            from .capcut.package_probe import run as probe_capcut
            capcut=probe_capcut(root)
            from .workflow_smoke_202 import run as workflow_202
            try:workflow=workflow_202(window,root)
            except Exception:
                # Keep the failed UI case visible, then collect independent
                # migration/admission/audio/export evidence. Never mark it PASS.
                workflow=dict(passed=False,error=traceback.format_exc())
            for page in (0,7,8,9,1,2,3):
                window.pages.setCurrentIndex(page);QApplication.processEvents()
                assert window.navigation_group.checkedId()==page
            from .release_smoke_203 import verify_admission
            admission_203=verify_admission(window,root,wait)
            automation=run_automation(window,root,wait)
            report={'passed':bool(workflow.get('passed')),'frozen':bool(getattr(sys,'frozen',False)),'version':VERSION,'platform':app.platformName(),'dpr':window.devicePixelRatioF(),
                    'tts':tts,'capcut':capcut,'automation':automation,'migration':migration,'gateway_recovery':recovery,'workflow_202':workflow,'admission_203':admission_203,
                    'txt_docx_fixture':True,'docx_export_units':2,'source_files_unchanged':original_hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (text,word)},
                    'themes':['light','dark'],'visible_pages':[0,7,8,9,1,2,3],'workspace_preserved':True,'queue_clear':True,'search_replace_undo':True,
                    'manual_popup':True,'auto_no_popup':True,'later_no_download':True,'download_verified':True,'real_http_requests':0,'installer_processes':0,
                    'isolated_database':str(store.path),'default_data_path_unchanged':True,'safe_exit_deadline_ms':900,'process_exit_observed_by_launcher':False}
            (root/'report.json').write_text(json.dumps(report,indent=2),encoding='utf8')
            window.confirm_exit=lambda running:'safe'
            window.close()
            QTimer.singleShot(900,lambda:verify_close())
        except Exception:
            (root/'failure.txt').write_text(traceback.format_exc(),encoding='utf8');window._exit_authorized=True;window.close();app.exit(1)
    def verify_close():
        if not window._closed:
            (root/'failure.txt').write_text('Safe Exit 900ms failed',encoding='utf8');app.exit(1)
        else:
            report=json.loads((root/'report.json').read_text(encoding='utf8'));report['safe_exit_verified']=True
            (root/'report.json').write_text(json.dumps(report,indent=2),encoding='utf8');app.quit()
    QTimer.singleShot(200,exercise)
    return app.exec()
